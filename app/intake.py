"""Client submissions (register, RP list, evidence) and "the system is notified of the upload" (02 §3.1-2).

submit() stores the bytes exactly as received, content-addressed under evidence/<sha256> with a server-computed
hash (a client-sent hash is never accepted). on_client_upload() runs in-process as the system principal: it freezes
the rows verbatim into audit.db, opens the engagement, runs D0 and D1/D2, writes the event and the UPLOAD anchor.
"""
import hashlib
import io
import json
import mimetypes
from pathlib import Path

from app import config, db

KINDS = ("SALES_REGISTER", "RP_LIST", "EVIDENCE")


def _next(c, table, prefix):
    r = c.execute(f"SELECT id FROM {table} WHERE id >= ? AND id < ? ORDER BY id DESC LIMIT 1",
                  (prefix + "0", prefix + ":")).fetchone()
    return f"{prefix}{(int(r[0][len(prefix):]) if r else 0) + 1:06d}"


def submit(kind, content, filename, submitted_by, request_id=None, declared=None):
    """Store bytes byte-for-byte (content-addressed, never overwritten) -> pbc.submissions id."""
    if kind not in KINDS:
        raise ValueError(f"INVALID_ARGUMENT: kind must be one of {', '.join(KINDS)}")
    if not isinstance(content, (bytes, bytearray)) or not content:
        raise ValueError("INVALID_ARGUMENT: the submission is empty")
    sha = hashlib.sha256(content).hexdigest()
    config.EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)
    path = config.EVIDENCE_DIR / sha
    if not path.exists():
        path.write_bytes(bytes(content))
    with db.borrow("pbc", "system") as p, db.tx(p):
        if request_id and not p.execute("SELECT 1 FROM evidence_requests WHERE id = ?", (request_id,)).fetchone():
            raise LookupError("NOT_FOUND: no evidence request with that id. List them with list_evidence_requests.")
        sid = _next(p, "submissions", "SUB-")
        p.execute("INSERT INTO submissions(id, engagement_id, kind, request_id, filename, mime, size, sha256, stored_path,"
                  " declared_json, submitted_by, submitted_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                  (sid, config.ENGAGEMENT_ID, kind, request_id, filename or "upload",
                   mimetypes.guess_type(filename or "")[0], len(content), sha, str(path),
                   json.dumps(declared, sort_keys=True, default=str) if declared is not None else None,
                   submitted_by, config.now()))
        if request_id:
            p.execute("UPDATE evidence_requests SET status = 'RECEIVED' WHERE id = ?", (request_id,))
    return sid


def ensure_engagement():
    """The FY engagement (03 #30), created with the config defaults on the first client upload."""
    eid = config.ENGAGEMENT_ID
    with db.borrow("audit", "system") as a, db.tx(a):
        if not a.execute("SELECT 1 FROM engagements WHERE id = ?", (eid,)).fetchone():
            a.execute("INSERT INTO engagements(id, entity_name, framework, period_start, period_end, status)"
                      " VALUES (?,?,?,?,?,'DRAFT')", (eid, config.ENTITY_NAME, config.FRAMEWORK, config.PERIOD_START,
                                                     config.PERIOD_END))
    return eid


def _load(submission_id):
    with db.borrow("pbc", "system") as p:
        s = p.execute("SELECT * FROM submissions WHERE id = ?", (submission_id,)).fetchone()
    if s is None:
        raise LookupError("NOT_FOUND: no submission with that id")
    content = Path(s["stored_path"]).read_bytes()
    if hashlib.sha256(content).hexdigest() != s["sha256"]:
        raise ValueError("PRECONDITION_FAILED: the stored file no longer matches its sha256 (evidence store tampered)")
    return s, content


def on_client_upload(submission_id):
    """System principal: freeze rows verbatim, open the engagement, run D0 + RP tagging, event + anchor."""
    from app import notify, register, related_parties, trail
    s, content = _load(submission_id)
    eid = ensure_engagement()
    out = {"engagement_id": eid, "submission_id": submission_id, "kind": s["kind"], "sha256": s["sha256"]}
    if s["kind"] == "EVIDENCE":
        notify.emit("EVIDENCE_RECEIVED", {"submission_id": submission_id, "request_id": s["request_id"]}, eid)
        return {**out, "next_steps": [{"tool": "test_sample", "example_args": {"run_id": "…", "sample_no": 1},
                                       "why": "test the sample this evidence answers"}]}
    headers, rows = register.parse_verbatim(content, s["filename"])
    with db.borrow("audit", "system") as a, db.tx(a):
        if s["kind"] == "SALES_REGISTER":
            colmap = register.map_columns(headers)
            a.executemany("INSERT INTO register_rows(engagement_id, submission_id, row_no, cells_json, mapped_json)"
                          " VALUES (?,?,?,?,?)",
                          [(eid, submission_id, r["row_no"], json.dumps(r["cells"], default=str),
                            json.dumps(register.map_row(headers, r, colmap), default=str)) for r in rows])
            a.execute("UPDATE engagements SET register_submission_id = ?, register_map_json = ? WHERE id = ?",
                      (submission_id, json.dumps({"headers": headers, "map": colmap}), eid))
        else:
            mapped = related_parties.map_rp_rows(headers, rows)
            by_row = {m["row_no"]: m for m in mapped}
            a.executemany("INSERT INTO rp_list_rows(engagement_id, submission_id, row_no, cells_json, mapped_json)"
                          " VALUES (?,?,?,?,?)",
                          [(eid, submission_id, r["row_no"], json.dumps(r["cells"], default=str),
                            json.dumps(by_row.get(r["row_no"]), default=str)) for r in rows if r["row_no"] in by_row])
            a.execute("UPDATE engagements SET rp_submission_id = ? WHERE id = ?", (submission_id, eid))
    if s["kind"] == "SALES_REGISTER":
        d0 = register.reconcile(eid)
        out.update(rows=len(rows), columns=len(headers), columns_mapped=sorted(colmap), d0_hits=len(d0))
    else:
        out.update(rows=len(mapped))
    rp_hits = related_parties.run_d1_d2(eid)
    out["rp_hits"] = len(rp_hits)
    out["anchor"] = trail.anchor(eid, "UPLOAD")["id"]
    notify.emit("REGISTER_UPLOADED" if s["kind"] == "SALES_REGISTER" else "RP_LIST_UPLOADED",
                {k: out[k] for k in out if k in ("submission_id", "rows", "d0_hits", "rp_hits", "sha256")}, eid)
    out["next_steps"] = [{"tool": "get_engagement_status", "example_args": {}, "why": "see what the upload surfaced"},
                         {"tool": "reconcile_register_to_system", "example_args": {},
                          "why": "every register vs system-of-record difference, row by row"}]
    return out

# ---------------------------------------------------------------- upload_sales_register(base_export_id, edits ...)

def build_register(base_export_id, edits=(), drop_rows=(), add_rows=(), role="accounts"):
    """A register built from one of the client's own exports, with edits applied: one call can produce a doctored
    8,000-row file (04 #26). edits [{row_no, column, value}] (row_no = 1-based data row, column = header text);
    drop_rows [row_no]; add_rows [{header: value}] appended. -> (bytes, filename)."""
    from openpyxl import load_workbook
    with db.borrow("pbc", role) as p:
        ex = p.execute("SELECT * FROM exports WHERE id = ?", (base_export_id,)).fetchone()
    if ex is None:
        raise LookupError("NOT_FOUND: no export with that id. Create one with export_sales_register.")
    wb = load_workbook(ex["stored_path"])
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(headers)}
    n = ws.max_row - 1
    for e in edits or ():
        r, h = e.get("row_no"), e.get("column")
        if not isinstance(r, int) or not 1 <= r <= n or h not in col:
            raise ValueError(f"INVALID_ARGUMENT: edit needs row_no 1..{n} and a column header from the export")
        cell = ws.cell(row=r + 1, column=col[h])
        cell.value = e.get("value")
        if isinstance(cell.value, str):
            cell.data_type = "s"
    for r in add_rows or ():
        ws.append([r.get(h) for h in headers])
    for r in sorted(set(drop_rows or ()), reverse=True):
        if not isinstance(r, int) or not 1 <= r <= n:
            raise ValueError(f"INVALID_ARGUMENT: drop_rows are data row numbers 1..{n}")
        ws.delete_rows(r + 1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"sales_register_{base_export_id}.xlsx"
