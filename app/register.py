"""The client's sales register: verbatim parse, column mapping, D0 register <-> system-of-record diff (05 D0).

parse_verbatim keeps every cell exactly as uploaded (raw value + number format, original column order); it
never cleans. map_row derives normalised fields used only to PAIR rows with SoR documents; the comparison
itself is on values, and every difference is a hit. Client text never enters a reason string: reasons cite
register row numbers and SoR ids, and the raw values go in evidence (shown under client_data by the tools).
"""
import csv
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

# ---------------------------------------------------------------- verbatim parse

def _norm_header(h):
    h = re.sub(r"\(.*?\)", "", str(h or "")).lower()
    return re.sub(r"[^a-z0-9]", "", h)     # 'Gross Amt.' -> 'grossamt'; 'Total Cost (Rs)' -> 'totalcost'

# canonical field -> normalised header spellings. Fig. 1's 21 headers first, then common ERP exports.
SYNONYMS = {
    "doc_no": ["billno", "invoiceno", "invno", "billnumber", "invoicenumber", "voucherno", "docno", "documentno",
               "bill", "invoice", "vchno"],
    "date": ["date", "invoicedate", "billdate", "invdate", "voucherdate", "docdate", "documentdate"],
    "job_no": ["jobno", "jobnumber", "job"],
    "customer": ["customer", "customername", "party", "partyname", "buyer", "buyername", "particulars",
                 "nameofcustomer", "nameofparty"],
    "ledger_account": ["ledgeraccount", "ledger", "salesledger", "salesaccount", "account"],
    "state_code": ["statecode", "state", "pos", "placeofsupply"],
    "gstin": ["gstno", "gstin", "gstinuin", "gstnumber", "customergstin", "buyergstin", "gstinno", "partygstin"],
    "taxable": ["totalcost", "taxablevalue", "taxable", "taxableamount", "taxableamt", "assessablevalue",
                "basicamount", "basicvalue"],
    "discount": ["discount", "disc", "discountamount", "discountamt"],
    "packing": ["packing", "packingcharges", "packingforwarding"],
    "freight": ["freight", "freightcharges"],
    "insurance": ["insurance", "insurancecharges"],
    "igst": ["igst", "igstamount", "igstamt", "integratedtax"],
    "cgst": ["cgst", "cgstamount", "cgstamt", "centraltax"],
    "sgst": ["sgst", "sgstamount", "sgstamt", "statetax", "utgst", "sgstutgst"],
    "tcs": ["tcscollection", "tcs", "tcsamount"],
    "round_off": ["round", "roundoff", "rounding", "roundedoff"],
    "gross": ["grossamt", "grossamount", "total", "invoicevalue", "invoicetotal", "totalamount", "grandtotal",
              "billamount", "invoiceamount"],
    "tax_desc": ["tax", "taxtype", "taxrate", "gstrate"],
    "bu": ["bu", "businessunit", "division"],
    "unit": ["unit", "series", "branch"],
    "einvoice": ["einvoice", "einvoicestatus", "irnstatus"],
    "irn": ["irn", "irnno", "irnnumber"],
    "ewb_no": ["ewaybill", "ewbno", "ewaybillno", "ewaybillnumber"],
    "hsn": ["hsn", "hsncode", "hsnsac", "sac"],
    "qty": ["qty", "quantity"],
}
_LOOKUP = {s: f for f, ss in SYNONYMS.items() for s in ss}


def map_columns(headers):
    """{canonical field: client header} for every header we recognise. Exact (normalised) match only:
    a wrong guess would silently mis-pair rows, so unknown headers stay unmapped. First occurrence wins."""
    out = {}
    for h in headers:
        f = _LOOKUP.get(_norm_header(h))
        if f and f not in out:
            out[f] = h
    return out


def _header_row(rows):
    """Index of the header row: the first of the top 15 rows that names >= 3 known columns incl. a doc no.
    Fig. 1 exports often carry title rows above the header."""
    for i, r in enumerate(rows[:15]):
        m = map_columns([c["v"] if isinstance(c["v"], str) else "" for c in r])
        if len(m) >= 3 and "doc_no" in m:
            return i
    return next((i for i, r in enumerate(rows) if any(c["v"] not in (None, "") for c in r)), 0)


def _cell_xlsx(c, cached):
    """One cell, verbatim: v = value (cached result for formulas), t = s|n|b|d|e|'' , fmt = number format."""
    v, fmt = cached.value, cached.number_format or "General"
    if isinstance(v, datetime):
        cell = {"v": v.isoformat(), "t": "d"}
    elif isinstance(v, date):
        cell = {"v": datetime(v.year, v.month, v.day).isoformat(), "t": "d"}
    elif isinstance(v, bool):
        cell = {"v": v, "t": "b"}
    elif isinstance(v, (int, float)):
        cell = {"v": v, "t": "n"}      # floats stay as Excel stored them; map_row converts via Decimal(repr)
    elif v is None:
        cell = {"v": None, "t": ""}
    else:
        cell = {"v": str(v), "t": "e" if cached.data_type == "e" else "s"}
    if fmt != "General":
        cell["fmt"] = fmt
    if c is not None and c.data_type == "f":
        cell["f"] = str(c.value)       # the formula text; listed in a Sheet 1 note (06 §2)
    return cell


def _rows_xlsx(content):
    from openpyxl import load_workbook
    wb_f = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    wb_v = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    best, best_score = None, -1
    for ws_f, ws_v in zip(wb_f.worksheets, wb_v.worksheets):
        if ws_v.sheet_state != "visible":
            continue
        rows = []
        for rf, rv in zip(ws_f.iter_rows(), ws_v.iter_rows()):
            rows.append([_cell_xlsx(cf, cv) for cf, cv in zip(rf, rv)])
        score = max((len(map_columns([c["v"] if isinstance(c["v"], str) else "" for c in r])) for r in rows[:15]),
                    default=0)
        if score > best_score:
            best, best_score = rows, score
    wb_f.close()
    wb_v.close()
    return best or []


def _rows_csv(content):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    # every CSV cell stays a string: '0012' keeps its zeros, '=HYPERLINK(...)' stays text
    return [[{"v": v, "t": "s", "fmt": "@"} for v in r] for r in csv.reader(io.StringIO(text, newline=""), dialect)]


def parse_verbatim(content, filename):
    """bytes -> (headers, rows). rows = [{"row_no": n, "cells": [cell per column]}], row_no = 1-based data row.
    Title rows above the header are dropped (06 §2 pastes header + data); trailing empty rows are dropped;
    nothing else is touched. Raises ValueError with an actionable message for unsupported files."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
        grid = _rows_xlsx(content)
    elif name.endswith((".csv", ".txt", ".tsv")):
        grid = _rows_csv(content)
    else:
        raise ValueError("INVALID_ARGUMENT: the register must be .xlsx or .csv (re-save .xls/.ods as .xlsx). "
                         "Example: upload_sales_register(content_base64=<xlsx bytes>, filename='sales_register.xlsx')")
    while grid and all(c["v"] in (None, "") for c in grid[-1]):
        grid.pop()
    if not grid:
        raise ValueError("INVALID_ARGUMENT: the register file has no rows. Export the full sales register and retry.")
    h = _header_row(grid)
    width = max(len(r) for r in grid[h:])
    headers = [str(c["v"]) if c["v"] is not None else "" for c in grid[h]] + [""] * (width - len(grid[h]))
    empty = {"v": None, "t": ""}
    rows = [{"row_no": i, "cells": r + [dict(empty) for _ in range(width - len(r))]}
            for i, r in enumerate(grid[h + 1:], 1)]
    return headers, rows

# ---------------------------------------------------------------- normalised fields (pairing only)

def doc_key(s):
    """Pairing key (05 D0): upper-case, split into letter/digit runs (drops spaces, '-', '/'), strip each
    number's leading zeros. 'G2/0012/25' == 'g2-12-25' == 'G2 12 25'; 'G2/1/225' stays different."""
    if s is None:
        return ""
    s = str(s).strip()
    if re.fullmatch(r"\d+\.0", s):     # an all-numeric bill no. that Excel stored as 12.0
        s = s[:-2]
    toks = re.findall(r"[A-Z]+|\d+", s.upper())
    return "|".join(t.lstrip("0") or "0" if t.isdigit() else t for t in toks)


def to_paise_signed(v):
    """Register amount -> signed paise for comparison only (raw value stays verbatim). Accepts numbers,
    '1,18,000.50', '(250000)', '-2,50,000', 'Rs 500'. Sub-paise float noise is rounded half-up.
    Returns None when it is not a number."""
    if v is None or v == "" or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        d = Decimal(repr(v))
    else:
        s = str(v).strip().replace("₹", "").replace(",", "").replace(" ", "")
        s = re.sub(r"^(rs\.?|inr)", "", s, flags=re.I)
        neg = s.startswith("(") and s.endswith(")")
        s = s.strip("()")
        try:
            d = Decimal(s)
        except InvalidOperation:
            return None
        if neg:
            d = -d
    if not d.is_finite():
        return None
    return int((d * 100).quantize(Decimal("1"), ROUND_HALF_UP))


_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct",
                                       "nov", "dec"], 1)}

def to_iso_date(v):
    """Register date -> 'YYYY-MM-DD' or None. Indian day-first: DD-MM-YYYY, DD/MM/YY, DD.MM.YYYY, 03-Apr-2025,
    ISO, xlsx datetimes ('d' cells) and Excel serials."""
    if v is None or v == "" or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if 20000 < v < 80000:          # Excel 1900-system serial
            return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
        return None
    s = str(v).strip()
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[T ][\d:.]+)?", s)
    if m:
        y, mo, d = int(m[1]), int(m[2]), int(m[3])
    else:
        m = re.fullmatch(r"(\d{1,2})[-/. ]([A-Za-z]{3,9}|\d{1,2})[-/. ](\d{2}|\d{4})", s)
        if not m:
            return None
        d, mo, y = int(m[1]), m[2], int(m[3])
        mo = _MONTHS.get(mo[:3].lower()) if mo.isalpha() else int(mo)
        if mo is None:
            return None
        if y < 100:
            y += 2000
    try:
        return date(y, mo, d).isoformat()
    except ValueError:
        return None


def fy_of(iso):
    """'2026-03-31' -> '2025-26' (FY derived from the date, never stored; 03 §1)."""
    d = date.fromisoformat(iso)
    start = d.year if d.month >= 4 else d.year - 1
    return f"{start}-{(start + 1) % 100:02d}"


def name_key(s):
    s = re.sub(r"[^A-Z0-9 ]", " ", str(s or "").upper().replace("M/S", " "))
    s = re.sub(r"\bPRIVATE\b", "PVT", re.sub(r"\bLIMITED\b", "LTD", s))
    return " ".join(s.split())


def gstin_key(s):
    return re.sub(r"\s", "", str(s or "")).upper()


def map_row(headers, row, colmap=None):
    """Normalised fields of one register row (audit.register_rows.mapped_json). Signed paise; a CN row
    keeps its negative sign. kind = ROW | TOTAL | BLANK; parse_errors names fields that did not parse."""
    colmap = colmap if colmap is not None else map_columns(headers)
    idx = {f: headers.index(h) for f, h in colmap.items()}
    val = lambda f: row["cells"][idx[f]]["v"] if f in idx else None
    raw_no = val("doc_no")
    texts = [c["v"] for c in row["cells"] if isinstance(c["v"], str)]
    if all(c["v"] in (None, "") for c in row["cells"]):
        kind = "BLANK"
    elif raw_no in (None, "") and any(re.match(r"\s*(grand\s*)?total\b", t, re.I) for t in texts):
        kind = "TOTAL"
    else:
        kind = "ROW"
    errors = []
    m = {"row_no": row["row_no"], "kind": kind, "doc_no": None if raw_no is None else str(raw_no),
         "doc_key": doc_key(raw_no), "customer": val("customer"), "gstin": val("gstin"),
         "state_code": val("state_code")}
    d = val("date")
    m["date"] = to_iso_date(d)
    if d not in (None, "") and m["date"] is None:
        errors.append("date")
    for f in ("taxable", "discount", "packing", "freight", "insurance", "igst", "cgst", "sgst", "tcs",
              "round_off", "gross"):
        v = val(f)
        p = to_paise_signed(v)
        m[f + "_paise"] = p
        if v not in (None, "") and p is None:
            errors.append(f)
    taxes = [m[f + "_paise"] for f in ("igst", "cgst", "sgst") if m[f + "_paise"] is not None]
    m["tax_paise"] = sum(taxes) if taxes else None
    m["is_cn"] = bool(re.match(r"\s*(CN|CRN|CDN)\b", str(raw_no or ""), re.I)) or (m["taxable_paise"] or 0) < 0
    m["parse_errors"] = errors
    m["cols"] = sorted(colmap)          # fields the client supplied; an absent column is never a difference
    return m

# ---------------------------------------------------------------- D0 (pure core; reconcile() feeds it from the DBs)

SEVERITY = {"NOT_IN_SOR": 5, "FY_DIFF": 5, "MISSING_FROM_REGISTER": 4, "DUPLICATE_ROW": 4, "CANCELLED_BUT_LISTED": 4,
            "CUSTOMER_DIFF": 4, "GSTIN_DIFF": 4, "TAXABLE_DIFF": 4, "TIEOUT_RESIDUE": 4, "DATE_DIFF": 3,
            "TAX_DIFF": 3, "TOTAL_DIFF": 3, "FORMAT_ONLY": 1}
STANDARDS = "SA 500 (inconsistent evidence from different sources must be resolved); SA 230"


def _inr(paise):
    """Indian grouping for server-written reasons. ponytail: local copy — switch to money.fmt_inr (IB-1)."""
    sign, p = ("-" if paise < 0 else ""), abs(paise)
    r, ps = divmod(p, 100)
    s = str(r)
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        head = ",".join([head[max(0, i - 2):i] for i in range(len(head), 0, -2)][::-1])
        s = head + "," + tail
    return f"{sign}Rs {s}.{ps:02d}"


def _hit(code, otype, oid, impact, reason, evidence):
    return {"detector": "D0", "code": code, "object_type": otype, "object_id": str(oid), "severity": SEVERITY[code],
            "impact_paise": abs(impact or 0), "reason": reason, "evidence": evidence}


def diff_register(mapped, sor_docs, period_start, period_end):
    """Every difference between register rows and SoR documents is a hit (05 D0).

    mapped:   map_row() dicts.
    sor_docs: [{id, doc_no, doc_type: INVOICE|CN, doc_date, party_id, party_name, gstin, taxable_paise,
                tax_paise, total_paise, status: ISSUED|CANCELLED}] with positive amounts (header sums of lines).
    Returns (hits, pairs) where pairs = {register row_no: sor id} for Sheet 1's SoR_Match / SoR_Doc_Id."""
    hits, pairs = [], {}
    by_key = {}
    for d in sor_docs:
        by_key.setdefault(doc_key(d["doc_no"]), d)
    seen = {}
    for m in mapped:
        if m["kind"] != "ROW":
            continue
        r = m["row_no"]
        ev = {"register_row": r, "register": {k: m[k] for k in ("doc_no", "date", "customer", "gstin",
                                                                "taxable_paise", "tax_paise", "gross_paise")}}
        if m["doc_key"] in seen and m["doc_key"]:
            hits.append(_hit("DUPLICATE_ROW", "register_row", r, m["taxable_paise"],
                             f"Register row {r} repeats the document already listed at row {seen[m['doc_key']]}; "
                             f"the duplicate adds {_inr(abs(m['taxable_paise'] or 0))} of taxable value.",
                             {**ev, "first_row": seen[m["doc_key"]]}))
            continue
        seen[m["doc_key"]] = r
        d = by_key.get(m["doc_key"]) if m["doc_key"] else None
        if d is None:
            hits.append(_hit("NOT_IN_SOR", "register_row", r, m["taxable_paise"],
                             f"Register row {r} has no matching invoice or credit note in the system of record "
                             f"({_inr(abs(m['taxable_paise'] or 0))} taxable is unsupported).", ev))
            continue
        pairs[r] = d["id"]
        ev["sor"] = {k: d.get(k) for k in ("id", "doc_no", "doc_date", "party_id", "party_name", "gstin",
                                           "taxable_paise", "tax_paise", "total_paise", "status")}
        hits += _compare(m, d, r, ev)
    listed = {m["doc_key"] for m in mapped if m["kind"] == "ROW"}
    for k, d in by_key.items():
        if (k not in listed and d["status"] == "ISSUED" and period_start <= d["doc_date"] <= period_end):
            what = "credit note" if d["doc_type"] == "CN" else "invoice"
            hits.append(_hit("MISSING_FROM_REGISTER", "credit_note" if d["doc_type"] == "CN" else "invoice", d["id"],
                             d["taxable_paise"],
                             f"{what.capitalize()} {d['id']} dated {d['doc_date']} is in the system of record but "
                             f"not in the uploaded register ({_inr(d['taxable_paise'])} taxable left out).",
                             {"sor": {k: d.get(k) for k in ("id", "doc_no", "doc_date", "party_id", "taxable_paise",
                                                            "tax_paise", "total_paise")}}))
    return hits, pairs


def _compare(m, d, r, ev):
    out, sid = [], d["id"]
    pre = f"Register row {r} vs {'credit note' if d['doc_type'] == 'CN' else 'invoice'} {sid}"
    if d["status"] == "CANCELLED":
        out.append(_hit("CANCELLED_BUT_LISTED", "register_row", r, m["taxable_paise"],
                        f"{pre}: the document is CANCELLED in the system of record but still listed in the register.", ev))
    if m["doc_no"] is not None and str(m["doc_no"]).strip() != d["doc_no"]:
        out.append(_hit("FORMAT_ONLY", "register_row", r, 0,
                        f"{pre}: the document number is written differently but matches after normalisation.", ev))
    cols = set(m["cols"])
    if "date" in cols and m["date"] != d["doc_date"]:
        out.append(_hit("DATE_DIFF", "register_row", r, m["taxable_paise"],
                        f"{pre}: the register date {m['date'] or '(unreadable)'} differs from the system date "
                        f"{d['doc_date']}.", ev))
        if m["date"] and fy_of(m["date"]) != fy_of(d["doc_date"]):
            out.append(_hit("FY_DIFF", "register_row", r, d["taxable_paise"],
                            f"{pre}: the register puts the document in FY {fy_of(m['date'])}, the system in FY "
                            f"{fy_of(d['doc_date'])}.", ev))
    if "customer" in cols and name_key(m["customer"]) != name_key(d["party_name"]):
        out.append(_hit("CUSTOMER_DIFF", "register_row", r, d["taxable_paise"],
                        f"{pre}: the customer in the register is not the system's party {d['party_id']}.", ev))
    elif "customer" in cols and str(m["customer"]).strip() != d["party_name"]:
        out.append(_hit("FORMAT_ONLY", "register_row", r, 0,
                        f"{pre}: the customer name is written differently but matches party {d['party_id']} "
                        f"after normalisation.", ev))
    if "gstin" in cols and gstin_key(m["gstin"]) != gstin_key(d.get("gstin")):
        out.append(_hit("GSTIN_DIFF", "register_row", r, d["taxable_paise"],
                        f"{pre}: the GSTIN in the register differs from the GSTIN on party {d['party_id']}"
                        f"{' (blank in the register)' if not gstin_key(m['gstin']) else ''}.", ev))
    sign = -1 if d["doc_type"] == "CN" else 1      # a CN may be listed negative (Fig. 1); compare magnitudes
    supplied = {"taxable": "taxable" in cols, "tax": bool(cols & {"igst", "cgst", "sgst"}), "gross": "gross" in cols}
    for f, code, label in (("taxable", "TAXABLE_DIFF", "taxable value"), ("tax", "TAX_DIFF", "tax"),
                           ("gross", "TOTAL_DIFF", "total")):
        if not supplied[f]:
            continue
        reg = m[f + "_paise"]
        if f == "tax" and reg is None:
            reg = 0                     # tax columns present but all blank: the register says no tax
        sor = d["total_paise" if f == "gross" else f + "_paise"]
        reg_mag = None if reg is None else (abs(reg) if sign < 0 else reg)
        if reg_mag != sor:
            out.append(_hit(code, "register_row", r, (reg_mag or 0) - sor,
                            f"{pre}: {label} {_inr(reg_mag) if reg_mag is not None else '(blank)'} in the register "
                            f"vs {_inr(sor)} in the system of record; difference {_inr((reg_mag or 0) - sor)}.", ev))
    return out


def tie_out(mapped, sor_docs, period_start, period_end, ledger_revenue_paise=None, gstr1_taxable_paise=None):
    """Totals tie-out (05 D0; Sheet 1 footer). Register net taxable (invoices - CNs) vs SoR vs ledger revenue vs
    GSTR-1. register - SoR is what the row-level D0 hits explain; the residue checks are the register's own
    total row vs the sum of its rows, and SoR vs ledger / GSTR-1. Returns (table, TIEOUT_RESIDUE hits)."""
    rows = [m for m in mapped if m["kind"] == "ROW"]
    reg = sum(-abs(m["taxable_paise"] or 0) if m["is_cn"] else (m["taxable_paise"] or 0) for m in rows)
    fy = [d for d in sor_docs if d["status"] == "ISSUED" and period_start <= d["doc_date"] <= period_end]
    sor = sum(-d["taxable_paise"] if d["doc_type"] == "CN" else d["taxable_paise"] for d in fy)
    footer = next((m["taxable_paise"] for m in mapped if m["kind"] == "TOTAL" and m["taxable_paise"] is not None), None)
    as_summed = sum(m["taxable_paise"] or 0 for m in rows)          # what Excel's SUM over the column gives
    table = {"register_rows": len(rows), "register_net_taxable_paise": reg, "register_footer_taxable_paise": footer,
             "sor_net_taxable_paise": sor, "register_minus_sor_paise": reg - sor,
             "ledger_revenue_paise": ledger_revenue_paise, "gstr1_taxable_paise": gstr1_taxable_paise}
    out = []
    if footer is not None and footer != as_summed:
        out.append(_hit("TIEOUT_RESIDUE", "engagement", "register_footer", footer - as_summed,
                        f"The register's own total row ({_inr(footer)}) does not equal the sum of its rows "
                        f"({_inr(as_summed)}); difference {_inr(footer - as_summed)}.", table))
    for key, label, other in (("ledger", "ledger revenue", ledger_revenue_paise),
                              ("gstr1", "GSTR-1 taxable", gstr1_taxable_paise)):
        if other is not None and other != sor:
            out.append(_hit("TIEOUT_RESIDUE", "engagement", key, sor - other,
                            f"System-of-record net taxable {_inr(sor)} vs {label} {_inr(other)}; difference "
                            f"{_inr(sor - other)}.", table))
    return table, out


def reconcile(engagement_id):
    """D0 over the frozen register (audit.register_rows) and erp.db. Wired once db.py lands (SH-2)."""
    raise NotImplementedError("register.reconcile: waiting for app.db (SH-2); the pure core is diff_register + tie_out")
