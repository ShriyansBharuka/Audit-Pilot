"""The client's system of record: write_doc + every business service (07 §4, 03 §2).

Every ERP write ends with write_doc in the same transaction: it appends the document's full after-image
(header + child rows, canonical JSON) to the hash-chained erp_log. After the transaction commits, the queued
events go to risk_graph.on_event (05 §19.6): best-effort, never blocking or rolling back the business write.

Services are (actor, role, **args) -> {"result": {...}, "next_steps": [...]}. Money is integer paise here (the MCP
tools convert rupee strings); dates are 'YYYY-MM-DD'; ids accept the internal id or the document number. Errors are
ValueError / PermissionError whose message starts with a 04 §3 code; the tool layer turns them into ToolError.
Constraints enforce integrity, never business policy (03 §1): an invoice without movement, a CN without inward
goods, a fake e-way bill number must all stay recordable so the engine can catch them.
"""
import base64
import hashlib
import json
import logging
import re
from contextlib import contextmanager
from datetime import date, datetime
from decimal import Decimal

from app import chain, config, db, ledger, money, portal

log = logging.getLogger("auditpilot.erp")

# doc_type -> (table, pk column, ((child table, fk column, order column), ...))
DOCS = {
    "party": ("parties", "id", ()),
    "party_link": ("party_links", "id", ()),
    "item": ("items", "id", ()),
    "scheme": ("schemes", "id", ()),
    "sales_order": ("sales_orders", "id", (("so_lines", "so_id", "line_no"),)),
    "movement": ("movements", "id", (("movement_lines", "movement_id", "item_id"),
                                     ("invoice_movements", "movement_id", "invoice_id"))),
    "invoice": ("invoices", "id", (("invoice_lines", "invoice_id", "line_no"),)),
    "credit_note": ("credit_notes", "id", (("cn_lines", "cn_id", "line_no"),)),
    "settlement": ("settlements", "id", ()),
    "purchase_bill": ("purchase_bills", "id", ()),
    "account": ("accounts", "code", ()),
    "journal": ("journals", "no", (("journal_lines", "journal_no", "line_no"),)),
}

# ---------------------------------------------------------------- the edit log

def image(conn, doc_type, doc_id):
    """The document as it is now: {"header": row, "<child table>": [rows]}; None if it doesn't exist."""
    table, pk, children = DOCS[doc_type]
    r = conn.execute(f"SELECT * FROM {table} WHERE {pk} = ?", (doc_id,)).fetchone()
    if r is None:
        return None
    out = {"header": dict(r)}
    for ctab, fk, order in children:
        out[ctab] = [dict(c) for c in conn.execute(f"SELECT * FROM {ctab} WHERE {fk} = ? ORDER BY {order}", (doc_id,))]
    return out


def _last_after(conn, doc_type, doc_id):
    r = conn.execute("SELECT after_json FROM erp_log WHERE doc_type = ? AND doc_id = ? ORDER BY seq DESC LIMIT 1",
                     (doc_type, str(doc_id))).fetchone()
    return r[0] if r else None


_PENDING = {}     # id(conn) -> [(doc_type, doc_id, seq)] waiting for the transaction to commit


def write_doc(conn, actor, role, tool, action, doc_type, doc_id, reason=None):
    """Append the doc's after-image to erp_log (03 §6). Call it after writing the typed rows, inside the
    transaction. before_json = the previous after_json of this doc. Returns the erp_log seq."""
    img = image(conn, doc_type, doc_id)
    if img is None:
        raise ValueError(f"NOT_FOUND: {doc_type} {doc_id} was not written before write_doc")
    seq, _h = chain.append(conn, "erp_log", {
        "recorded_at": config.now(), "actor": actor, "role": role, "tool": tool, "action": action,
        "doc_type": doc_type, "doc_id": str(doc_id), "before_json": _last_after(conn, doc_type, doc_id),
        "after_json": chain.canon(img), "reason": reason})
    _PENDING.setdefault(id(conn), []).append((doc_type, str(doc_id), seq))
    return seq


def flush_events(conn):
    """After COMMIT: hand each logged document to AuditGraph. Errors are logged, never raised."""
    from app import risk_graph
    out = []
    for doc_type, doc_id, seq in _PENDING.pop(id(conn), []):
        try:
            out += risk_graph.on_event(doc_type, doc_id, seq) or []
        except Exception:  # noqa: BLE001 - best-effort by contract (05 §19.6)
            log.exception("risk_graph.on_event failed for %s %s (seq %s)", doc_type, doc_id, seq)
    return out


def discard_events(conn):
    _PENDING.pop(id(conn), None)

# ---------------------------------------------------------------- integrity helpers (used by evidence.verify_integrity)

def all_images(conn, doc_types=None):
    """{(doc_type, doc_id): image} for every document, built in one pass per table (same shape as image())."""
    out = {}
    for doc_type, (table, pk, children) in DOCS.items():
        if doc_types and doc_type not in doc_types:
            continue
        imgs = {}
        for r in conn.execute(f"SELECT * FROM {table}"):
            imgs[r[pk]] = {"header": dict(r), **{ctab: [] for ctab, _f, _o in children}}
        for ctab, fk, order in children:
            for r in conn.execute(f"SELECT * FROM {ctab} ORDER BY {fk}, {order}"):
                if r[fk] in imgs:
                    imgs[r[fk]][ctab].append(dict(r))
        for k, v in imgs.items():
            out[(doc_type, str(k))] = v
    return out


def out_of_band(conn, docs=None):
    """Documents whose current rows differ from their latest erp_log after-image (03 §6 check 2): an edit made
    outside the application (direct SQLite). kind = CHANGED | MISSING (logged, now gone) | UNLOGGED (no log row).
    docs: optional set of (doc_type, doc_id) to check (a sample's bundle); default every document."""
    latest = {}
    for r in conn.execute("SELECT doc_type, doc_id, seq, after_json, actor FROM erp_log WHERE seq IN "
                          "(SELECT MAX(seq) FROM erp_log GROUP BY doc_type, doc_id)"):
        if docs is None or (r[0], r[1]) in docs:
            latest[(r[0], r[1])] = (r[2], r[3], r[4])
    current = all_images(conn, {d for d, _i in docs} if docs is not None else None)
    out = []
    for (doc_type, doc_id), now in current.items():
        if docs is not None and (doc_type, doc_id) not in docs:
            continue
        got = latest.pop((doc_type, doc_id), None)
        if got is None:
            out.append({"doc_type": doc_type, "doc_id": doc_id, "kind": "UNLOGGED", "seq": None})
            continue
        seq, after, actor = got
        if chain.canon(now) != after:
            out.append({"doc_type": doc_type, "doc_id": doc_id, "kind": "CHANGED", "seq": seq,
                        "last_logged_by": actor, "fields": _diff(json.loads(after), now)})
    for (doc_type, doc_id), (seq, _a, actor) in latest.items():
        out.append({"doc_type": doc_type, "doc_id": doc_id, "kind": "MISSING", "seq": seq, "last_logged_by": actor})
    return out


def _diff(before, after):
    """{'header.lr_date': [old, new], 'invoice_lines[1].qty': [old, new]} between two images."""
    out = {}
    for k in sorted(set(before) | set(after)):
        b, a = before.get(k), after.get(k)
        if isinstance(b, dict) or isinstance(a, dict):
            b, a = b or {}, a or {}
            for f in sorted(set(b) | set(a)):
                if b.get(f) != a.get(f):
                    out[f"{k}.{f}"] = [b.get(f), a.get(f)]
        elif b != a:
            for i in range(max(len(b or []), len(a or []))):
                rb = (b or [])[i] if i < len(b or []) else None
                ra = (a or [])[i] if i < len(a or []) else None
                if rb != ra:
                    for f in sorted(set(rb or {}) | set(ra or {})):
                        if (rb or {}).get(f) != (ra or {}).get(f):
                            out[f"{k}[{i}].{f}"] = [(rb or {}).get(f), (ra or {}).get(f)]
    return out

# ---------------------------------------------------------------- service plumbing

SERVICES = ("create_party", "link_parties", "create_item", "create_lead", "create_scheme", "create_sales_order",
            "issue_invoice", "generate_eway_bill", "record_dispatch", "record_delivery", "issue_credit_note",
            "record_receipt", "record_payment", "record_purchase", "post_journal_entry", "file_gstr1", "file_gstr3b",
            "amend_record", "cancel_document", "import_records", "export_sales_register", "render_document")

TAGS = {"SELF", "CUSTOMER", "VENDOR", "TRANSPORTER", "DIRECTOR", "KMP", "RELATIVE"}
keep_open = db.keep_open          # bulk loads (the seeder, replay): one connection per (db, role)


@contextmanager
def _session(role):
    """One erp transaction for one service call; AuditGraph events fire only after COMMIT."""
    with db.borrow("erp", role) as c:
        ok = False
        try:
            with db.tx(c):
                yield c
            ok = True
        finally:
            if ok:
                flush_events(c)
            else:
                discard_events(c)


def _ok(result, *steps):
    return {"result": result, "next_steps": [{"tool": t, "example_args": a, "why": w} for t, a, w in steps]}


def _bad(msg):
    raise ValueError(f"INVALID_ARGUMENT: {msg}")


def _date(v, name, default=None):
    v = v if v not in (None, "") else default
    if not isinstance(v, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
        _bad(f"{name} must be a date 'YYYY-MM-DD'. Example: {name}='2026-03-30'")
    try:
        date.fromisoformat(v)
    except ValueError:
        _bad(f"{name} is not a real calendar date")
    return v


def _int(v, name, minimum=0, default=None):
    v = default if v is None else v
    if isinstance(v, bool) or not isinstance(v, int) or v < minimum:
        _bad(f"{name} must be an integer >= {minimum} (money inside services is integer paise)")
    return v


def _one_of(v, name, allowed):
    if v not in allowed:
        _bad(f"{name} must be one of {', '.join(allowed)}")
    return v


def _next_id(c, table, col, prefix, width):
    """Next zero-padded id after the highest numeric one (a primary-key range seek; 'P-DIR1' style ids are skipped)."""
    r = c.execute(f"SELECT {col} FROM {table} WHERE {col} >= ? AND {col} < ? ORDER BY {col} DESC LIMIT 1",
                  (prefix + "0", prefix + ":")).fetchone()
    n = int(re.match(r"\d*", r[0][len(prefix):]).group() or 0) if r else 0
    return f"{prefix}{n + 1:0{width}d}"


def fy_start(d):
    y, m = int(d[:4]), int(d[5:7])
    return f"{y if m >= 4 else y - 1}-04-01"


def _next_no(c, series, d):
    fs = fy_start(d)
    c.execute("INSERT OR IGNORE INTO doc_counters(series, fy_start, next_no) VALUES (?, ?, 1)", (series, fs))
    n = c.execute("SELECT next_no FROM doc_counters WHERE series = ? AND fy_start = ?", (series, fs)).fetchone()[0]
    c.execute("UPDATE doc_counters SET next_no = next_no + 1 WHERE series = ? AND fy_start = ?", (series, fs))
    return n, int(fs[2:4])


_REFS = {"invoice": ("invoices", "id", "invoice_no"), "credit_note": ("credit_notes", "id", "cn_no"),
         "sales_order": ("sales_orders", "id", "so_no"), "movement": ("movements", "id", "lr_no"),
         "party": ("parties", "id", None), "item": ("items", "id", None), "scheme": ("schemes", "id", None),
         "purchase_bill": ("purchase_bills", "id", "bill_no")}
_FIND = {"invoice": "query_records(entity='invoices')", "credit_note": "query_records(entity='credit_notes')",
         "sales_order": "query_records(entity='sales_orders')", "movement": "query_records(entity='movements')",
         "party": "query_records(entity='parties')", "item": "query_records(entity='items')",
         "scheme": "query_records(entity='schemes')", "purchase_bill": "query_records(entity='purchase_bills')"}


def _get(c, kind, ref):
    """The row for an id or document number; NOT_FOUND names the tool that lists valid ids (no client text)."""
    table, pk, no = _REFS[kind]
    r = c.execute(f"SELECT * FROM {table} WHERE {pk} = ?", (ref,)).fetchone()
    if r is None and no:
        r = c.execute(f"SELECT * FROM {table} WHERE {no} = ? ORDER BY {pk} LIMIT 1", (ref,)).fetchone()
    if r is None:
        raise LookupError(f"NOT_FOUND: no {kind.replace('_', ' ')} with that id or number. List them with {_FIND[kind]}.")
    return r


def _unit(series):
    if series not in config.UNITS:
        _bad(f"series must be a unit: {', '.join(config.UNITS)}")
    return config.UNITS[series]


def _self_account(c):
    r = c.execute("SELECT bank_account FROM parties WHERE id = 'SELF'").fetchone()
    return (r[0] if r and r[0] else "KMCL-CA-0001")


def _is_service(hsn):
    return str(hsn or "").startswith("99")


def _revenue_account(hsn, pos):
    return "SALES_SERVICE" if _is_service(hsn) else ("SALES_EXPORT" if pos == "96" else "SALES")


def _norm_lines(c, lines, with_rate=True):
    if not isinstance(lines, list) or not lines:
        _bad("lines must be a non-empty list. Example: lines=[{item_id:'IT-001', qty:10, rate_paise:680000}]")
    out = []
    for i, ln in enumerate(lines, 1):
        if not isinstance(ln, dict):
            _bad(f"line {i} must be an object")
        it = _get(c, "item", ln.get("item_id"))
        row = {"item_id": it["id"], "qty": _int(ln.get("qty"), f"line {i} qty")}
        if with_rate:
            row["rate_paise"] = _int(ln.get("rate_paise"), f"line {i} rate_paise", default=it["list_rate_paise"])
            row["discount_paise"] = _int(ln.get("discount_paise"), f"line {i} discount_paise", default=0)
            if ln.get("scheme_id"):
                row["scheme_id"] = _get(c, "scheme", ln["scheme_id"])["id"]
            if ln.get("gst_rate_bp") is not None:
                row["gst_rate_bp"] = _int(ln["gst_rate_bp"], f"line {i} gst_rate_bp")
        out.append(row)
    return out

# ---------------------------------------------------------------- masters

def create_party(actor, role, id=None, kind="ENTITY", tags="CUSTOMER", name=None, pan=None, gstin=None,
                 state_code=None, address="", pincode="", phone=None, email=None, bank_account=None, ifsc=None):
    """Customer, vendor, transporter, director, KMP or relative. Sales may create customers only."""
    if not name:
        _bad("name is required. Example: create_party(name='Arabian Shipyards Pvt Ltd', gstin='27AAPFU0939F1ZV')")
    _one_of(kind, "kind", ("ENTITY", "PERSON"))
    tag_list = [t.strip().upper() for t in str(tags or "").split(",") if t.strip()]
    if not tag_list or any(t not in TAGS for t in tag_list) or ("SELF" in tag_list and id != "SELF"):
        _bad(f"tags is a comma list of {', '.join(sorted(TAGS - {'SELF'}))}")
    if role == "sales" and tag_list != ["CUSTOMER"]:
        raise PermissionError("PERMISSION_DENIED: sales may create customers only (tags='CUSTOMER'); vendors, "
                              "transporters, directors and relatives are created by client_admin")
    with _session(role) as c:
        pid = id or _next_id(c, "parties", "id", "P-", 4)
        if c.execute("SELECT 1 FROM parties WHERE id = ?", (pid,)).fetchone():
            raise ValueError("CONFLICT: a party with that id already exists")
        if not state_code and gstin and len(str(gstin).strip()) >= 2:
            state_code = str(gstin).strip()[:2]
        c.execute("INSERT INTO parties(id, kind, tags, name, pan, gstin, state_code, address, pincode, phone, email,"
                  " bank_account, ifsc, recorded_by, recorded_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                  (pid, kind, ",".join(tag_list), name, pan, gstin, state_code, address or "", pincode or "", phone,
                   email, bank_account, ifsc, actor, config.now()))
        write_doc(c, actor, role, "create_party", "CREATE", "party", pid)
    return _ok({"id": pid}, ("create_lead", {"party_id": pid, "enquired_on": config.today(), "channel": "EMAIL"},
                             "record the first customer interaction"))


def link_parties(actor, role, from_party=None, to_party=None, relation=None, effective_from=None, effective_to=None,
                 pct_bp=None, source="MBP-1", id=None):
    """Director / KMP / relative / shareholding / control link with dates (MBP-1 style): the RP truth."""
    _one_of(relation, "relation", ("DIRECTOR_OF", "KMP_OF", "SHAREHOLDER_OF", "RELATIVE_OF", "CONTROLS", "PARTNER_OF"))
    _one_of(source, "source", ("MBP-1", "DIR-12", "REGISTER", "OTHER"))
    effective_from = _date(effective_from, "effective_from")
    if effective_to:
        effective_to = _date(effective_to, "effective_to")
    with _session(role) as c:
        a, b = _get(c, "party", from_party)["id"], _get(c, "party", to_party)["id"]
        lid = id or _next_id(c, "party_links", "id", "L-", 3)
        c.execute("INSERT INTO party_links(id, from_party, to_party, relation, pct_bp, effective_from, effective_to,"
                  " source) VALUES (?,?,?,?,?,?,?,?)", (lid, a, b, relation, pct_bp, effective_from, effective_to, source))
        write_doc(c, actor, role, "link_parties", "CREATE", "party_link", lid)
    return _ok({"id": lid})


def create_item(actor, role, name=None, hsn=None, uom="NOS", gst_rate_bp=1800, list_rate_paise=0, cost_paise=0, id=None):
    if not name or not hsn:
        _bad("name and hsn are required. Example: create_item(name='Ball valve DN25', hsn='8481', gst_rate_bp=1800)")
    with _session(role) as c:
        iid = id or _next_id(c, "items", "id", "IT-", 3)
        c.execute("INSERT INTO items(id, name, hsn, uom, gst_rate_bp, list_rate_paise, cost_paise) VALUES (?,?,?,?,?,?,?)",
                  (iid, name, str(hsn), uom, _int(gst_rate_bp, "gst_rate_bp"), _int(list_rate_paise, "list_rate_paise"),
                   _int(cost_paise, "cost_paise")))
        write_doc(c, actor, role, "create_item", "CREATE", "item", iid)
    return _ok({"id": iid})


def create_scheme(actor, role, name=None, kind="TRADE_DISCOUNT", rate_bp=0, valid_from=None, valid_to=None,
                  announced_on=None, approved_by=None, min_qty=None, id=None):
    """A discount / rebate scheme (the promotions that drive demand)."""
    if not name:
        _bad("name is required. Example: create_scheme(name='March push', kind='TRADE_DISCOUNT', rate_bp=1200, "
             "valid_from='2026-03-22', valid_to='2026-03-31')")
    _one_of(kind, "kind", ("TRADE_DISCOUNT", "VOLUME_REBATE", "CASH_DISCOUNT", "FREE_QTY"))
    valid_from, valid_to = _date(valid_from, "valid_from"), _date(valid_to, "valid_to")
    announced_on = _date(announced_on, "announced_on", default=config.today())
    with _session(role) as c:
        sid = id or _next_id(c, "schemes", "id", "SCH-", 3)
        c.execute("INSERT INTO schemes(id, name, kind, rate_bp, min_qty, valid_from, valid_to, announced_on, approved_by)"
                  " VALUES (?,?,?,?,?,?,?,?,?)", (sid, name, kind, _int(rate_bp, "rate_bp"), min_qty, valid_from,
                                                  valid_to, announced_on, approved_by))
        write_doc(c, actor, role, "create_scheme", "CREATE", "scheme", sid)
    return _ok({"id": sid})

# ---------------------------------------------------------------- order to cash

_CHANNELS = ("EMAIL", "PHONE", "VISIT", "PORTAL", "SHOPIFY")
_TERMS = ("EX_WORKS", "FOR_DISPATCH", "FOR_DESTINATION", "BILL_AND_HOLD", "CONSIGNMENT")


def create_lead(actor, role, party_id=None, enquired_on=None, channel="EMAIL", contact_name=None, contact_phone=None,
                contact_email=None, id=None):
    """First customer interaction: a sales order in ENQUIRY."""
    _one_of(channel, "channel", _CHANNELS)
    enquired_on = _date(enquired_on, "enquired_on", default=config.today())
    with _session(role) as c:
        pid = _get(c, "party", party_id)["id"]
        sid = id or _next_id(c, "sales_orders", "id", "SO-", 6)
        c.execute("INSERT INTO sales_orders(id, so_no, party_id, status, enquired_on, channel, contact_name, contact_phone,"
                  " contact_email) VALUES (?,?,?,'ENQUIRY',?,?,?,?,?)",
                  (sid, sid, pid, enquired_on, channel, contact_name, contact_phone, contact_email))
        write_doc(c, actor, role, "create_lead", "CREATE", "sales_order", sid)
    return _ok({"id": sid, "status": "ENQUIRY"},
               ("create_sales_order", {"so_id": sid, "quoted_on": enquired_on, "lines": [{"item_id": "IT-001", "qty": 1}]},
                "quote / confirm the order with prices"))


def create_sales_order(actor, role, lines=None, party_id=None, so_id=None, channel=None, enquired_on=None,
                       quoted_on=None, confirmed_on=None, contact_name=None, contact_phone=None, contact_email=None,
                       customer_po_no=None, delivery_terms="FOR_DISPATCH", credit_days=30, scheme_id=None,
                       bill_and_hold_json=None, id=None):
    """Commitment and pricing. With so_id it converts a lead; otherwise it creates the SO in one step."""
    _one_of(delivery_terms, "delivery_terms", _TERMS)
    for k, v in (("quoted_on", quoted_on), ("confirmed_on", confirmed_on), ("enquired_on", enquired_on)):
        if v:
            _date(v, k)
    status = "CONFIRMED" if confirmed_on else ("QUOTED" if quoted_on else "ENQUIRY")
    bh = json.dumps(bill_and_hold_json, sort_keys=True) if isinstance(bill_and_hold_json, dict) else bill_and_hold_json
    with _session(role) as c:
        norm = _norm_lines(c, lines)
        sch = _get(c, "scheme", scheme_id)["id"] if scheme_id else None
        if so_id:
            so = _get(c, "sales_order", so_id)
            if c.execute("SELECT 1 FROM so_lines WHERE so_id = ?", (so["id"],)).fetchone():
                raise ValueError("CONFLICT: this sales order already has lines; use amend_record(entity='sales_orders')")
            sid, action = so["id"], "AMEND"
            c.execute("UPDATE sales_orders SET status=?, quoted_on=?, confirmed_on=?, customer_po_no=?, delivery_terms=?,"
                      " credit_days=?, scheme_id=?, bill_and_hold_json=? WHERE id=?",
                      (status, quoted_on, confirmed_on, customer_po_no, delivery_terms, credit_days, sch, bh, sid))
        else:
            if not channel:
                channel = "EMAIL"
            _one_of(channel, "channel", _CHANNELS)
            pid = _get(c, "party", party_id)["id"]
            sid, action = id or _next_id(c, "sales_orders", "id", "SO-", 6), "CREATE"
            c.execute("INSERT INTO sales_orders(id, so_no, party_id, status, enquired_on, channel, contact_name,"
                      " contact_phone, contact_email, quoted_on, confirmed_on, customer_po_no, delivery_terms, credit_days,"
                      " scheme_id, bill_and_hold_json) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                      (sid, sid, pid, status, enquired_on or quoted_on or confirmed_on or config.today(), channel,
                       contact_name, contact_phone, contact_email, quoted_on, confirmed_on, customer_po_no, delivery_terms,
                       credit_days, sch, bh))
        for i, ln in enumerate(norm, 1):
            c.execute("INSERT INTO so_lines(so_id, line_no, item_id, qty, rate_paise, discount_paise) VALUES (?,?,?,?,?,?)",
                      (sid, i, ln["item_id"], ln["qty"], ln["rate_paise"], ln["discount_paise"]))
        write_doc(c, actor, role, "create_sales_order", action, "sales_order", sid)
    return _ok({"id": sid, "status": status},
               ("issue_invoice", {"so_id": sid, "series": "G2", "invoice_date": config.today()}, "bill the order"))


def issue_invoice(actor, role, series="G2", invoice_date=None, so_id=None, party_id=None, place_of_supply=None,
                  lines=None, buyer_gstin=None, irn=None, register_irn=True, tax_mode="AUTO", round_off_paise=None,
                  invoice_no=None, id=None):
    """Tax invoice. The number is system-assigned per unit per FY (Rule 46(b)); `invoice_no` is honoured only for
    import_records. Registers the IRN for B2B (unless the client records its own `irn` or register_irn=False) and
    posts the SALES voucher. tax_mode IGST | CGST_SGST overrides the place-of-supply rule (plantable)."""
    unit = _unit(series)
    invoice_date = _date(invoice_date, "invoice_date", default=config.today())
    _one_of(tax_mode, "tax_mode", ("AUTO", "IGST", "CGST_SGST"))
    with _session(role) as c:
        so = _get(c, "sales_order", so_id) if so_id else None
        party = _get(c, "party", party_id or (so["party_id"] if so else None))
        if lines is None and so is not None:
            lines = [dict(r) for r in c.execute("SELECT item_id, qty, rate_paise, discount_paise FROM so_lines"
                                                " WHERE so_id = ? ORDER BY line_no", (so["id"],))]
        norm = _norm_lines(c, lines)
        pos = str(place_of_supply or party["state_code"] or unit["state_code"])
        inter = (pos != unit["state_code"]) if tax_mode == "AUTO" else (tax_mode == "IGST")
        b_gstin = buyer_gstin if buyer_gstin is not None else party["gstin"]
        rows, taxable, tax = [], 0, 0
        for i, ln in enumerate(norm, 1):
            it = c.execute("SELECT hsn, gst_rate_bp FROM items WHERE id = ?", (ln["item_id"],)).fetchone()
            t = ln["qty"] * ln["rate_paise"] - ln["discount_paise"]
            if t < 0:
                _bad(f"line {i}: discount exceeds the line value")
            rate = ln.get("gst_rate_bp", it["gst_rate_bp"])
            tx = money.tax_split(t, rate, inter)
            rows.append({**ln, "line_no": i, "hsn": it["hsn"], "taxable_paise": t, "gst_rate_bp": rate, **tx})
            taxable += t
            tax += sum(tx.values())
        ro = money.round_off(taxable + tax) if round_off_paise is None else int(round_off_paise)
        if invoice_no:
            if c.execute("SELECT 1 FROM invoices WHERE invoice_no = ?", (invoice_no,)).fetchone():
                raise ValueError("CONFLICT: an invoice with that number already exists")
            no = invoice_no
        else:
            n, yy = _next_no(c, series, invoice_date)
            no = f"{series}/{n:04d}/{yy:02d}"
        iid = id or _next_id(c, "invoices", "id", "INV-", 6)
        ack = None
        if irn is None and register_irn and b_gstin and config.AATO_EINVOICE:
            ack = portal.register_irn(unit["gstin"], portal.fy_label(invoice_date), "INV", no, invoice_date, b_gstin,
                                      taxable, tax)
            irn = ack["irn"]
        c.execute("INSERT INTO invoices(id, invoice_no, series, supplier_gstin, invoice_date, party_id, so_id, buyer_gstin,"
                  " place_of_supply, irn, round_off_paise) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
                  (iid, no, series, unit["gstin"], invoice_date, party["id"], so["id"] if so else None, b_gstin, pos,
                   irn, ro))
        for r in rows:
            c.execute("INSERT INTO invoice_lines(invoice_id, line_no, item_id, hsn, qty, rate_paise, discount_paise,"
                      " scheme_id, taxable_paise, gst_rate_bp, igst_paise, cgst_paise, sgst_paise)"
                      " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
                      (iid, r["line_no"], r["item_id"], r["hsn"], r["qty"], r["rate_paise"], r["discount_paise"],
                       r.get("scheme_id"), r["taxable_paise"], r["gst_rate_bp"], r["igst_paise"], r["cgst_paise"],
                       r["sgst_paise"]))
        write_doc(c, actor, role, "issue_invoice", "CREATE", "invoice", iid)
        jl = [{"account_code": "DEBTORS", "party_id": party["id"], "dr_paise": taxable + tax + ro}] if taxable + tax + ro > 0 else []
        by_acct = {}
        for r in rows:
            a = _revenue_account(r["hsn"], pos)
            by_acct[a] = by_acct.get(a, 0) + r["taxable_paise"]
        jl += [{"account_code": a, "cr_paise": v} for a, v in sorted(by_acct.items()) if v]
        for code, k in (("OUTPUT_IGST", "igst_paise"), ("OUTPUT_CGST", "cgst_paise"), ("OUTPUT_SGST", "sgst_paise")):
            v = sum(r[k] for r in rows)
            if v:
                jl.append({"account_code": code, "cr_paise": v})
        if ro:
            jl.append({"account_code": "ROUND_OFF", "cr_paise": ro} if ro > 0 else {"account_code": "ROUND_OFF", "dr_paise": -ro})
        jno = ledger.post_journal(c, invoice_date, "SALES", jl, f"Sales invoice {no}", actor, role, "issue_invoice",
                                  source_type="invoice", source_id=iid) if taxable + tax > 0 else None
    goods = any(not _is_service(r["hsn"]) and not r["item_id"].startswith("CHG") for r in rows)
    steps = [("generate_eway_bill", {"invoice_id": iid, "vehicle_no": "MH08AB1234", "distance_km": 120},
              "goods over Rs 50,000 need an e-way bill before they move")] if goods else []
    steps.append(("record_dispatch" if goods else "record_receipt",
                  {"invoice_nos": [no], "moved_on": invoice_date} if goods else {"invoice_nos": [no]},
                  "move the goods" if goods else "record the customer's payment"))
    return _ok({"id": iid, "invoice_no": no, "irn": irn, "taxable_paise": taxable, "tax_paise": tax,
                "round_off_paise": ro, "total_paise": taxable + tax + ro, "journal_no": jno}, *steps)


def _invoice_totals(c, iid):
    r = c.execute("SELECT SUM(taxable_paise), SUM(igst_paise + cgst_paise + sgst_paise) FROM invoice_lines"
                  " WHERE invoice_id = ?", (iid,)).fetchone()
    return (r[0] or 0), (r[1] or 0)


def generate_eway_bill(actor, role, invoice_id=None, vehicle_no=None, distance_km=None, reason="SUPPLY",
                       transport_doc_no=None, value_override_paise=None):
    """Register an e-way bill on the mock portal for an invoice. value_override_paise plants a value mismatch."""
    _one_of(reason, "reason", ("SUPPLY", "JOB_WORK", "SKD", "LINE_SALES", "OWN_USE", "EXHIBITION", "RETURN", "OTHERS"))
    distance_km = _int(distance_km, "distance_km")
    with db.borrow("erp", role) as c:
        inv = _get(c, "invoice", invoice_id)
        party = _get(c, "party", inv["party_id"])
        taxable, tax = _invoice_totals(c, inv["id"])
        hsn = c.execute("SELECT hsn FROM invoice_lines WHERE invoice_id = ? AND hsn NOT LIKE '99%' ORDER BY line_no LIMIT 1",
                        (inv["id"],)).fetchone()
    value = taxable + tax if value_override_paise is None else _int(value_override_paise, "value_override_paise")
    ewb = portal.generate_eway_bill(
        supplier_gstin=inv["supplier_gstin"], doc_type="INV", doc_no=inv["invoice_no"], doc_date=inv["invoice_date"],
        to_gstin=inv["buyer_gstin"], from_pincode=config.UNITS[inv["series"]]["pincode"],
        to_pincode=party["pincode"] or "000000", distance_km=distance_km, value_paise=value, hsn=hsn[0] if hsn else None,
        vehicle_no=vehicle_no, transport_doc_no=transport_doc_no, reason=reason)
    return _ok({**ewb, "invoice_id": inv["id"], "value_paise": value},
               ("record_dispatch", {"invoice_nos": [inv["invoice_no"]], "ewb_no": ewb["ewb_no"], "vehicle_no": vehicle_no},
                "record the movement with this e-way bill"))


def record_dispatch(actor, role, moved_on=None, invoice_nos=None, transporter_id=None, vehicle_no=None, lr_no=None,
                    lr_date=None, from_pincode=None, to_pincode=None, lines=None, direction="OUT", ewb_no=None,
                    party_id=None, so_id=None, gate_pass_no=None, id=None):
    """A physical movement (OUT = dispatch, IN = sales-return inward). ewb_no is free text as the client typed it."""
    _one_of(direction, "direction", ("OUT", "IN"))
    moved_on = _date(moved_on, "moved_on", default=config.today())
    if lr_date:
        lr_date = _date(lr_date, "lr_date")
    with _session(role) as c:
        invs = [_get(c, "invoice", n) for n in (invoice_nos or [])]
        if direction == "OUT" and not invs and not party_id:
            _bad("a dispatch needs invoice_nos (or party_id for goods sent without an invoice)")
        party = _get(c, "party", party_id or (invs[0]["party_id"] if invs else None))
        so = _get(c, "sales_order", so_id)["id"] if so_id else (invs[0]["so_id"] if invs else None)
        if lines is None:
            agg = {}
            for inv in invs:
                for r in c.execute("SELECT item_id, qty, hsn FROM invoice_lines WHERE invoice_id = ?", (inv["id"],)):
                    if not _is_service(r["hsn"]) and not r["item_id"].startswith("CHG"):
                        agg[r["item_id"]] = agg.get(r["item_id"], 0) + r["qty"]
            norm = [{"item_id": k, "qty": v} for k, v in agg.items()]
        else:
            agg = {}
            for ln in _norm_lines(c, lines, with_rate=False):
                agg[ln["item_id"]] = agg.get(ln["item_id"], 0) + ln["qty"]
            norm = [{"item_id": k, "qty": v} for k, v in agg.items()]
        if transporter_id:
            transporter_id = _get(c, "party", transporter_id)["id"]
        unit_pin = config.UNITS[invs[0]["series"]]["pincode"] if invs else config.UNITS["G2"]["pincode"]
        mid = id or _next_id(c, "movements", "id", "MV-", 6)
        c.execute("INSERT INTO movements(id, direction, party_id, so_id, moved_on, from_pincode, to_pincode, transporter_id,"
                  " vehicle_no, lr_no, lr_date, gate_pass_no) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                  (mid, direction, party["id"], so, moved_on,
                   from_pincode or (unit_pin if direction == "OUT" else party["pincode"]),
                   to_pincode or (party["pincode"] if direction == "OUT" else unit_pin),
                   transporter_id, vehicle_no, lr_no, lr_date or (moved_on if lr_no else None), gate_pass_no))
        for ln in norm:
            c.execute("INSERT INTO movement_lines(movement_id, item_id, qty) VALUES (?,?,?)", (mid, ln["item_id"], ln["qty"]))
        for inv in invs:
            c.execute("INSERT INTO invoice_movements(invoice_id, movement_id, ewb_no) VALUES (?,?,?)",
                      (inv["id"], mid, None if ewb_no is None else str(ewb_no)))
        write_doc(c, actor, role, "record_dispatch", "CREATE", "movement", mid)
    return _ok({"id": mid, "invoices": [i["id"] for i in invs], "lines": len(norm)},
               ("record_delivery", {"movement_id": mid, "delivered_on": moved_on, "received_by": "stores"},
                "record proof of delivery (transfer of control for FOR-destination terms)"))


def record_delivery(actor, role, movement_id=None, delivered_on=None, received_by=None):
    """Proof of delivery: an amendment of the movement (its [V] history stays in erp_log)."""
    delivered_on = _date(delivered_on, "delivered_on", default=config.today())
    with _session(role) as c:
        m = _get(c, "movement", movement_id)
        c.execute("UPDATE movements SET delivered_on = ?, received_by = ? WHERE id = ?", (delivered_on, received_by, m["id"]))
        write_doc(c, actor, role, "record_delivery", "AMEND", "movement", m["id"], reason="POD")
    return _ok({"id": m["id"], "delivered_on": delivered_on})


_CN_ACCOUNT = {"SALES_RETURN": "SALES_RETURNS", "RATE_DIFF": "DISCOUNTS_ALLOWED", "SCHEME_REBATE": "DISCOUNTS_ALLOWED",
               "QUALITY": "SALES_RETURNS", "OTHER": "SALES_RETURNS"}


def issue_credit_note(actor, role, lines=None, cn_date=None, reason="OTHER", reason_text=None, scheme_id=None,
                      return_movement_id=None, party_id=None, series=None, register_irn=True, id=None):
    """Credit note over one or more invoices (s.34(1)); lines [{invoice_no?, item_id?, qty?, taxable_paise}]. Tax follows
    the credited invoice line's rate and supply type. Posts CREDIT_NOTE and registers the CRN IRN for B2B."""
    _one_of(reason, "reason", tuple(_CN_ACCOUNT))
    cn_date = _date(cn_date, "cn_date", default=config.today())
    if not isinstance(lines, list) or not lines:
        _bad("lines is required. Example: lines=[{invoice_no:'G2/0391/25', taxable_paise:1559322}]")
    with _session(role) as c:
        rows, party, unit_s = [], None, series
        for i, ln in enumerate(lines, 1):
            inv = _get(c, "invoice", ln["invoice_no"]) if ln.get("invoice_no") else None
            it = _get(c, "item", ln["item_id"]) if ln.get("item_id") else None
            t = _int(ln.get("taxable_paise"), f"line {i} taxable_paise")
            p = inv["party_id"] if inv else party_id
            if party and p and p != party:
                _bad("all lines of a credit note must be for one customer")
            party = party or p
            rate, inter = 1800, None
            if inv:
                unit_s = unit_s or inv["series"]
                q = c.execute("SELECT gst_rate_bp, igst_paise FROM invoice_lines WHERE invoice_id = ? AND (? IS NULL OR"
                              " item_id = ?) ORDER BY line_no LIMIT 1", (inv["id"], it["id"] if it else None,
                                                                        it["id"] if it else None)).fetchone()
                if q:
                    rate, inter = q["gst_rate_bp"], q["igst_paise"] > 0
                else:
                    inter = inv["place_of_supply"] != config.UNITS[inv["series"]]["state_code"]
            elif it:
                rate = it["gst_rate_bp"]
            rows.append({"invoice_id": inv["id"] if inv else None, "item_id": it["id"] if it else None,
                         "qty": ln.get("qty"), "taxable_paise": t, "rate": rate, "inter": inter})
        pty = _get(c, "party", party or party_id)
        unit_s = unit_s or "G2"
        unit = _unit(unit_s)
        for r in rows:
            if r["inter"] is None:
                r["inter"] = (pty["state_code"] or unit["state_code"]) != unit["state_code"]
            r.update(money.tax_split(r["taxable_paise"], r["rate"], r["inter"]))
        taxable = sum(r["taxable_paise"] for r in rows)
        tax = sum(r["igst_paise"] + r["cgst_paise"] + r["sgst_paise"] for r in rows)
        n, yy = _next_no(c, f"CN-{unit_s}", cn_date)
        no = f"CN/{unit_s}/{n:04d}/{yy:02d}"
        cid = id or _next_id(c, "credit_notes", "id", "CN-", 6)
        irn = None
        if register_irn and pty["gstin"] and config.AATO_EINVOICE:
            irn = portal.register_irn(unit["gstin"], portal.fy_label(cn_date), "CRN", no, cn_date, pty["gstin"], taxable,
                                      tax)["irn"]
        c.execute("INSERT INTO credit_notes(id, cn_no, cn_date, party_id, reason, reason_text, scheme_id,"
                  " return_movement_id, irn) VALUES (?,?,?,?,?,?,?,?,?)",
                  (cid, no, cn_date, pty["id"], reason, reason_text, _get(c, "scheme", scheme_id)["id"] if scheme_id else None,
                   _get(c, "movement", return_movement_id)["id"] if return_movement_id else None, irn))
        for i, r in enumerate(rows, 1):
            c.execute("INSERT INTO cn_lines(cn_id, line_no, invoice_id, item_id, qty, taxable_paise, igst_paise, cgst_paise,"
                      " sgst_paise) VALUES (?,?,?,?,?,?,?,?,?)", (cid, i, r["invoice_id"], r["item_id"], r["qty"],
                                                                 r["taxable_paise"], r["igst_paise"], r["cgst_paise"],
                                                                 r["sgst_paise"]))
        write_doc(c, actor, role, "issue_credit_note", "CREATE", "credit_note", cid)
        jl = [{"account_code": _CN_ACCOUNT[reason], "dr_paise": taxable}] if taxable else []
        for code, k in (("OUTPUT_IGST", "igst_paise"), ("OUTPUT_CGST", "cgst_paise"), ("OUTPUT_SGST", "sgst_paise")):
            v = sum(r[k] for r in rows)
            if v:
                jl.append({"account_code": code, "dr_paise": v})
        jl.append({"account_code": "DEBTORS", "party_id": pty["id"], "cr_paise": taxable + tax})
        jno = ledger.post_journal(c, cn_date, "CREDIT_NOTE", jl, f"Credit note {no}", actor, role, "issue_credit_note",
                                  source_type="credit_note", source_id=cid) if taxable + tax > 0 else None
    return _ok({"id": cid, "cn_no": no, "irn": irn, "taxable_paise": taxable, "tax_paise": tax,
                "total_paise": taxable + tax, "journal_no": jno})


def _outstanding(c, inv_id):
    t, x = _invoice_totals(c, inv_id)
    ro = c.execute("SELECT round_off_paise FROM invoices WHERE id = ?", (inv_id,)).fetchone()[0]
    paid = c.execute("SELECT COALESCE(SUM(amount_paise), 0) FROM settlements WHERE direction = 'RECEIPT'"
                     " AND doc_type = 'INVOICE' AND doc_id = ?", (inv_id,)).fetchone()[0]
    return t + x + ro - paid


def record_receipt(actor, role, amount_paise=None, invoice_nos=None, settled_on=None, payer_name=None,
                   payer_account=None, payer_ifsc=None, utr=None, mode="NEFT", party_id=None, narration=None):
    """A customer payment: one bank statement line (ext) + settlements against invoices + a RECEIPT voucher.
    The payer's account is recorded as the bank shows it (D4 THIRD_PARTY_PAYER / PAYER_IS_RP compare it)."""
    amount = _int(amount_paise, "amount_paise", minimum=1)
    settled_on = _date(settled_on, "settled_on", default=config.today())
    with _session(role) as c:
        invs = [_get(c, "invoice", n) for n in (invoice_nos or [])]
        party = _get(c, "party", party_id or (invs[0]["party_id"] if invs else None))
        bt = portal.bank_txn(_self_account(c), settled_on, amount, payer_name or party["name"],
                             payer_account if payer_account is not None else party["bank_account"],
                             payer_ifsc if payer_ifsc is not None else party["ifsc"], utr, narration or "receipt")
        left, settled = amount, []
        for inv in invs:
            take = min(left, max(0, _outstanding(c, inv["id"])))
            if take <= 0:
                continue
            settled.append(("INVOICE", inv["id"], take))
            left -= take
        if left > 0:
            settled.append(("ADVANCE", None, left))
        ids = []
        for doc_type, doc_id, amt in settled:
            sid = _next_id(c, "settlements", "id", "ST-", 6)
            c.execute("INSERT INTO settlements(id, direction, settled_on, party_id, doc_type, doc_id, amount_paise, bank_txn_id,"
                      " mode) VALUES (?,?,?,?,?,?,?,?,?)", (sid, "RECEIPT", settled_on, party["id"], doc_type, doc_id, amt,
                                                            bt, mode))
            write_doc(c, actor, role, "record_receipt", "CREATE", "settlement", sid)
            ids.append(sid)
        jno = ledger.post_journal(c, settled_on, "RECEIPT", [{"account_code": "BANK", "dr_paise": amount},
                                  {"account_code": "DEBTORS", "party_id": party["id"], "cr_paise": amount}],
                                  f"Receipt {bt}", actor, role, "record_receipt", source_type="bank_txn", source_id=bt)
    return _ok({"bank_txn_id": bt, "settlements": ids, "journal_no": jno, "unallocated_paise": left})


def record_purchase(actor, role, vendor_id=None, bill_no=None, bill_date=None, taxable_paise=None, tax_paise=0,
                    item_id=None, qty=None, id=None):
    """A purchase bill (T2; closes goods round trips for D10). Posts PURCHASE."""
    bill_date = _date(bill_date, "bill_date", default=config.today())
    if not bill_no:
        _bad("bill_no is required (the vendor's bill number)")
    t, x = _int(taxable_paise, "taxable_paise", minimum=1), _int(tax_paise, "tax_paise")
    with _session(role) as c:
        v = _get(c, "party", vendor_id)
        it = _get(c, "item", item_id)["id"] if item_id else None
        pid = id or _next_id(c, "purchase_bills", "id", "PB-", 6)
        c.execute("INSERT INTO purchase_bills(id, vendor_id, bill_no, bill_date, item_id, qty, taxable_paise, tax_paise)"
                  " VALUES (?,?,?,?,?,?,?,?)", (pid, v["id"], bill_no, bill_date, it, qty, t, x))
        write_doc(c, actor, role, "record_purchase", "CREATE", "purchase_bill", pid)
        jl = [{"account_code": "PURCHASES", "dr_paise": t}] + ([{"account_code": "INPUT_IGST", "dr_paise": x}] if x else [])
        jl.append({"account_code": "CREDITORS", "party_id": v["id"], "cr_paise": t + x})
        jno = ledger.post_journal(c, bill_date, "PURCHASE", jl, f"Purchase bill {pid}", actor, role, "record_purchase",
                                  source_type="purchase_bill", source_id=pid)
    return _ok({"id": pid, "journal_no": jno})


def record_payment(actor, role, amount_paise=None, vendor_id=None, settled_on=None, payee_account=None,
                   payee_ifsc=None, utr=None, mode="NEFT", purchase_ids=None, narration=None):
    """A vendor payment (T2; closes cash round trips for D10): bank line out + settlement + PAYMENT voucher."""
    amount = _int(amount_paise, "amount_paise", minimum=1)
    settled_on = _date(settled_on, "settled_on", default=config.today())
    with _session(role) as c:
        v = _get(c, "party", vendor_id)
        bt = portal.bank_txn(_self_account(c), settled_on, -amount, v["name"],
                             payee_account if payee_account is not None else v["bank_account"],
                             payee_ifsc if payee_ifsc is not None else v["ifsc"], utr, narration or "payment")
        bills = [_get(c, "purchase_bill", b)["id"] for b in (purchase_ids or [])]
        sid = _next_id(c, "settlements", "id", "ST-", 6)
        c.execute("INSERT INTO settlements(id, direction, settled_on, party_id, doc_type, doc_id, amount_paise, bank_txn_id,"
                  " mode) VALUES (?,?,?,?,?,?,?,?,?)", (sid, "PAYMENT", settled_on, v["id"],
                                                        "PURCHASE" if bills else "ADVANCE", bills[0] if bills else None,
                                                        amount, bt, mode))
        write_doc(c, actor, role, "record_payment", "CREATE", "settlement", sid)
        jno = ledger.post_journal(c, settled_on, "PAYMENT", [{"account_code": "CREDITORS", "party_id": v["id"],
                                                              "dr_paise": amount}, {"account_code": "BANK", "cr_paise": amount}],
                                  f"Payment {bt}", actor, role, "record_payment", source_type="bank_txn", source_id=bt)
    return _ok({"bank_txn_id": bt, "settlement_id": sid, "journal_no": jno})


def post_journal_entry(actor, role, voucher_date=None, lines=None, narration=None, aje_ref=None):
    """A balanced manual JV (narration mandatory). Management also posts an accepted AJE this way (aje_ref)."""
    voucher_date = _date(voucher_date, "voucher_date", default=config.today())
    with _session(role) as c:
        jno = ledger.post_journal(c, voucher_date, "JOURNAL", lines or [], narration, actor, role, "post_journal_entry",
                                  source_type="manual", source_id=None, aje_ref=aje_ref)
    return _ok({"journal_no": jno})

# ---------------------------------------------------------------- GST returns

def _gstins(gstin):
    if gstin:
        if gstin not in {u["gstin"] for u in config.UNITS.values()}:
            _bad(f"gstin must be a company registration: {', '.join(u['gstin'] for u in config.UNITS.values())}")
        return [gstin]
    return [u["gstin"] for u in config.UNITS.values()]


def file_gstr1(actor, role, period=None, gstin=None, omit_doc_nos=(), changes=()):
    """File a month's GSTR-1 from the books for one or both registrations. omit_doc_nos / changes make the filing
    differ from the books (plantable, D6)."""
    def unit_of(doc_no):              # 'G2/0001/25' -> G2, 'CN/G11/0003/25' -> G11
        parts = str(doc_no or "").split("/")
        return parts[1] if parts[0] == "CN" and len(parts) > 1 else parts[0]

    out = []
    for g in _gstins(gstin):
        u = next(k for k, v in config.UNITS.items() if v["gstin"] == g)
        mine = lambda no: bool(gstin) or unit_of(no) == u
        try:
            out.append(portal.file_gstr1(g, period, omit=[n for n in omit_doc_nos or () if mine(n)],
                                         changes=[ch for ch in changes or () if mine(ch.get("doc_no"))]))
        except portal.Refused as e:
            raise ValueError(str(e)) from e
    return _ok({"filings": out})


def file_gstr3b(actor, role, period=None, gstin=None, figures=None):
    try:
        return _ok({"filings": [portal.file_gstr3b(g, period, figures) for g in _gstins(gstin)]})
    except portal.Refused as e:
        raise ValueError(str(e)) from e

# ---------------------------------------------------------------- amend / cancel / import

_AMENDABLE = {"parties": "party", "party_links": "party_link", "items": "item", "schemes": "scheme",
              "sales_orders": "sales_order", "movements": "movement", "invoice_movements": "movement", "accounts": "account"}
_FINANCIAL = {"invoices": "invoice", "invoice_lines": "invoice", "credit_notes": "credit_note", "cn_lines": "credit_note",
              "settlements": "settlement", "purchase_bills": "purchase bill", "journals": "journal",
              "journal_lines": "journal"}
_OWNER = {"sales": {"parties", "sales_orders", "schemes"}, "warehouse": {"movements", "invoice_movements"}}


def amend_record(actor, role, entity=None, id=None, changes=None, reason=None):
    """Change a [V] record with a reason; before/after are kept in erp_log. Financial documents are refused."""
    if entity in _FINANCIAL:
        if entity in ("invoices", "invoice_lines", "credit_notes", "cn_lines"):
            raise ValueError("PRECONDITION_FAILED: invoices and credit notes are financial documents and are never edited. "
                             "Cancel and re-issue with cancel_document(doc_type='invoice', id=...) then issue_invoice(...), "
                             "or reduce the sale with issue_credit_note(...).")
        raise ValueError("PRECONDITION_FAILED: ledger and settlement records are never edited; post a reversal or a new "
                         "voucher with post_journal_entry(...).")
    if entity not in _AMENDABLE:
        _bad(f"entity must be one of {', '.join(_AMENDABLE)}")
    if role != "client_admin" and entity not in _OWNER.get(role, set()):
        raise PermissionError(f"PERMISSION_DENIED: role {role} may not amend {entity}; owners: sales (parties, "
                              "sales_orders, schemes), warehouse (movements, invoice_movements), client_admin (all)")
    if not isinstance(changes, dict) or not changes:
        _bad("changes must be an object of column: new value. Example: changes={'lr_date': '2026-04-02'}")
    if not (reason or "").strip():
        _bad("reason is required for every amendment (it is kept in the edit log)")
    doc_type = _AMENDABLE[entity]
    table, pk, _c = DOCS[doc_type] if entity != "invoice_movements" else ("invoice_movements", None, ())
    with _session(role) as c:
        cols = {r[1] for r in c.execute(f"PRAGMA table_info({entity})")}
        bad = [k for k in changes if k not in cols or k in ("id", "code", "recorded_by", "recorded_at", "invoice_id",
                                                            "movement_id")]
        if bad:
            _bad(f"these columns can't be changed on {entity}: {', '.join(map(str, bad))}")
        sets = ", ".join(f"{k} = ?" for k in changes)
        if entity == "invoice_movements":
            inv_id, _, mov_id = str(id).partition("|")
            if not mov_id:
                _bad("id for invoice_movements is '<invoice_id>|<movement_id>'")
            cur = c.execute(f"UPDATE invoice_movements SET {sets} WHERE invoice_id = ? AND movement_id = ?",
                            (*changes.values(), inv_id, mov_id))
            doc_id = mov_id
        else:
            row = c.execute(f"SELECT * FROM {table} WHERE {pk} = ?", (id,)).fetchone()
            if row is None and entity == "sales_orders":
                row = c.execute("SELECT * FROM sales_orders WHERE so_no = ?", (id,)).fetchone()
            if row is None:
                raise LookupError(f"NOT_FOUND: no {entity} row with that id. List them with query_records(entity='{entity}').")
            if role == "sales" and entity == "parties" and "CUSTOMER" not in (row["tags"] or "").split(","):
                raise PermissionError("PERMISSION_DENIED: sales may amend customers only")
            doc_id = row[pk]
            cur = c.execute(f"UPDATE {table} SET {sets} WHERE {pk} = ?", (*changes.values(), doc_id))
        if cur.rowcount == 0:
            raise LookupError("NOT_FOUND: nothing to amend with that id")
        write_doc(c, actor, role, "amend_record", "AMEND", doc_type, doc_id, reason=reason)
    return _ok({"entity": entity, "id": str(doc_id), "changed": sorted(changes)})


def cancel_document(actor, role, doc_type=None, id=None, reason=None, cancelled_on=None):
    """Cancel an invoice or credit note: status CANCELLED + a REVERSAL journal; asks the IRP to cancel the IRN,
    which it refuses after 24 h (the books still cancel; the portal keeps the IRN active)."""
    _one_of(doc_type, "doc_type", ("invoice", "credit_note"))
    if not (reason or "").strip():
        _bad("reason is required")
    cancelled_on = _date(cancelled_on, "cancelled_on", default=config.today())
    with _session(role) as c:
        d = _get(c, doc_type, id)
        if d["status"] == "CANCELLED":
            raise ValueError("CONFLICT: the document is already cancelled")
        if doc_type == "invoice":
            c.execute("UPDATE invoices SET status = 'CANCELLED', cancelled_on = ? WHERE id = ?", (cancelled_on, d["id"]))
        else:
            c.execute("UPDATE credit_notes SET status = 'CANCELLED' WHERE id = ?", (d["id"],))
        write_doc(c, actor, role, "cancel_document", "CANCEL", doc_type, d["id"], reason=reason)
        reversed_ = []
        for j in c.execute("SELECT no FROM journals WHERE source_type = ? AND source_id = ? AND voucher_type IN"
                           " ('SALES', 'CREDIT_NOTE') AND no NOT IN (SELECT reverses_no FROM journals WHERE reverses_no"
                           " IS NOT NULL)", (doc_type, d["id"])).fetchall():
            reversed_.append(ledger.reverse_journal(c, j[0], cancelled_on, actor, role, "cancel_document", reason))
        irn_status = None
        if d["irn"]:
            try:
                irn_status = portal.cancel_irn(d["irn"])["status"]
            except portal.Refused as e:
                irn_status = "REFUSED: " + str(e).split(": ", 1)[-1]
    return _ok({"id": d["id"], "status": "CANCELLED", "reversal_journals": reversed_, "irn": irn_status})


_IMPORTABLE = {"parties": create_party, "party_links": link_parties, "items": create_item, "schemes": create_scheme,
               "leads": create_lead, "sales_orders": create_sales_order, "invoices": issue_invoice,
               "credit_notes": issue_credit_note, "receipts": record_receipt, "purchases": record_purchase}


def import_records(actor, role, entity=None, rows=None):
    """Bulk create (<= 1,000 rows) through the same services, one transaction per row; per-row errors.
    Imported invoices keep their source numbers (and are tested the same way)."""
    if entity not in _IMPORTABLE:
        _bad(f"entity must be one of {', '.join(_IMPORTABLE)}")
    if not isinstance(rows, list) or not rows:
        _bad("rows must be a non-empty list of objects with the same fields the single-record tool takes")
    if len(rows) > 1000:
        _bad("at most 1,000 rows per call; split the batch")
    created, errors = [], []
    for i, row in enumerate(rows, 1):
        try:
            res = _IMPORTABLE[entity](actor, role, **row)
            r = res["result"]
            created.append(r.get("id") or r.get("bank_txn_id"))
        except TypeError:
            errors.append({"row": i, "error": "INVALID_ARGUMENT: unknown field for this entity"})
        except (ValueError, LookupError, PermissionError) as e:
            errors.append({"row": i, "error": str(e)})
    return _ok({"entity": entity, "created": len(created), "ids": created[:50], "errors": errors[:50],
                "error_count": len(errors)})

# ---------------------------------------------------------------- register export (Fig. 1) and documents

FIG1 = ["Bill no.", "Date", "Job No.", "Customer", "Ledger Account", "State Code", "GST No.", "Total Cost", "Discount",
        "Packing", "Freight", "Insurance", "IGST", "CGST", "SGST", "TCSCollection", "Round", "Gross Amt.", "Tax", "BU",
        "E Invoice"]
_BU = (("8481", "Valves"), ("8412", "Actuators"), ("8501", "Actuators"), ("853", "Panels"), ("902", "Instruments"),
       ("8544", "Cables"), ("99", "Services"))


def _bu(hsn):
    return next((b for p, b in _BU if str(hsn or "").startswith(p)), "Marine equipment")


def register_rows(c, fy_start_iso="2025-04-01"):
    """The register as the client's ERP would export it (Fig. 1 columns, CNs as negative rows), as python values."""
    fy_end = f"{int(fy_start_iso[:4]) + 1}-03-31"
    rows = []
    for inv in c.execute("SELECT i.*, p.name AS pname, s.so_no FROM invoices i JOIN parties p ON p.id = i.party_id"
                         " LEFT JOIN sales_orders s ON s.id = i.so_id WHERE i.status = 'ISSUED' AND i.invoice_date"
                         " BETWEEN ? AND ? ORDER BY i.invoice_date, i.invoice_no", (fy_start_iso, fy_end)).fetchall():
        ls = c.execute("SELECT * FROM invoice_lines WHERE invoice_id = ? ORDER BY line_no", (inv["id"],)).fetchall()
        t = sum(r["taxable_paise"] for r in ls)
        ig, cg, sg = (sum(r[k] for r in ls) for k in ("igst_paise", "cgst_paise", "sgst_paise"))
        chg = {k: sum(r["taxable_paise"] for r in ls if r["item_id"] == f"CHG-{k}") for k in ("PACKING", "FREIGHT", "INSURANCE")}
        rate = ls[0]["gst_rate_bp"] if ls else 0
        hsn0 = ls[0]["hsn"] if ls else ""
        ledger_acct = ("SALES ACCOUNT- SERVICE" if _is_service(hsn0) else
                       ("EXPORT SALES" if inv["place_of_supply"] == "96" else "SALES ACCOUNT"))
        tax_txt = f"IGST {rate // 100}%" if ig else f"CGST/SGST {rate // 200}%"
        rows.append([inv["invoice_no"], inv["invoice_date"], inv["so_no"] or "", inv["pname"], ledger_acct,
                     int(inv["place_of_supply"]) if str(inv["place_of_supply"]).isdigit() else 0, inv["buyer_gstin"] or "",
                     t, sum(r["discount_paise"] for r in ls), chg["PACKING"], chg["FREIGHT"], chg["INSURANCE"], ig, cg, sg,
                     0, inv["round_off_paise"], t + ig + cg + sg + inv["round_off_paise"], tax_txt, _bu(hsn0),
                     "Complete" if inv["irn"] else "Incomplete"])
    for cn in c.execute("SELECT n.*, p.name AS pname, p.gstin AS pgstin, p.state_code FROM credit_notes n JOIN parties p"
                        " ON p.id = n.party_id WHERE n.status = 'ISSUED' AND n.cn_date BETWEEN ? AND ?"
                        " ORDER BY n.cn_date, n.cn_no", (fy_start_iso, fy_end)).fetchall():
        ls = c.execute("SELECT * FROM cn_lines WHERE cn_id = ?", (cn["id"],)).fetchall()
        t = sum(r["taxable_paise"] for r in ls)
        ig, cg, sg = (sum(r[k] for r in ls) for k in ("igst_paise", "cgst_paise", "sgst_paise"))
        rows.append([cn["cn_no"], cn["cn_date"], "", cn["pname"], "SALES RETURN",
                     int(cn["state_code"]) if str(cn["state_code"] or "").isdigit() else 0, cn["pgstin"] or "",
                     -t, 0, 0, 0, 0, -ig, -cg, -sg, 0, 0, -(t + ig + cg + sg), "IGST" if ig else "CGST/SGST", "",
                     "Complete" if cn["irn"] else "Incomplete"])
    rows.sort(key=lambda r: (r[1], r[0]))
    return rows


def export_sales_register(actor, role, fy_start="2025-04-01"):
    """Build the client's sales register (Fig. 1 columns) from erp.db as an .xlsx export -> {export_id, sha256}."""
    import io

    from openpyxl import Workbook
    fy_start = _date(fy_start, "fy_start")
    c = db.connect("erp", role)
    try:
        rows = register_rows(c, fy_start)
    finally:
        c.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales Register"
    ws.append(FIG1)
    money_cols = range(7, 18)
    for rn, r in enumerate(rows, start=2):    # ws.max_row rescans every cell: O(n^2) on 7k rows (115 s)
        out = list(r)
        out[1] = datetime.fromisoformat(r[1])
        for k in money_cols:
            out[k] = Decimal(r[k]) / 100
        ws.append(out)
        ws.cell(row=rn, column=2).number_format = "DD-MM-YYYY"
        for k in money_cols:
            ws.cell(row=rn, column=k + 1).number_format = "#,##,##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    sha = hashlib.sha256(data).hexdigest()
    config.EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    path = config.EXPORTS_DIR / f"{sha}.xlsx"
    if not path.exists():
        path.write_bytes(data)
    p = db.connect("pbc", role)
    try:
        with db.tx(p):
            eid = _next_id(p, "exports", "id", "EXP-", 4)
            p.execute("INSERT INTO exports(id, kind, fy_start, sha256, stored_path, created_by, created_at)"
                      " VALUES (?,?,?,?,?,?,?)", (eid, "SALES_REGISTER", fy_start, sha, str(path), actor, config.now()))
    finally:
        p.close()
    return _ok({"export_id": eid, "sha256": sha, "rows": len(rows)},
               ("upload_sales_register", {"base_export_id": eid}, "send the register to the auditor (edits optional)"))


def _kv(pairs):
    return "\n".join(f"{k}: {'' if v is None else v}" for k, v in pairs) + "\n"


def render_document(actor, role, doc_type=None, id=None, format="text"):
    """The official text (or PDF) of an INVOICE, EWAY_BILL, LR, POD or CREDIT_NOTE built from the books + portal, as
    'Key: Value' lines, so an agent can answer an evidence request (or tamper with a copy)."""
    _one_of(doc_type, "doc_type", ("INVOICE", "EWAY_BILL", "LR", "POD", "CREDIT_NOTE"))
    _one_of(format, "format", ("text", "pdf"))
    c = db.connect("erp", role)
    x = db.connect("ext", role)
    try:
        if doc_type == "INVOICE":
            inv = _get(c, "invoice", id)
            p = _get(c, "party", inv["party_id"])
            ls = c.execute("SELECT l.*, i.name FROM invoice_lines l JOIN items i ON i.id = l.item_id WHERE invoice_id = ?"
                           " ORDER BY line_no", (inv["id"],)).fetchall()
            t, tx = _invoice_totals(c, inv["id"])
            pairs = [("Document", "TAX INVOICE"), ("Invoice No", inv["invoice_no"]), ("Invoice Date", inv["invoice_date"]),
                     ("Supplier", config.ENTITY_NAME), ("Supplier GSTIN", inv["supplier_gstin"]), ("Buyer", p["name"]),
                     ("Buyer GSTIN", inv["buyer_gstin"]), ("Place of Supply", inv["place_of_supply"]), ("IRN", inv["irn"])]
            pairs += [(f"Line {r['line_no']}", f"{r['item_id']} | {r['name']} | HSN {r['hsn']} | qty {r['qty']} | "
                                               f"taxable {money.to_rupees(r['taxable_paise'])}") for r in ls]
            pairs += [("Taxable Value", money.to_rupees(t)), ("IGST", money.to_rupees(sum(r["igst_paise"] for r in ls))),
                      ("CGST", money.to_rupees(sum(r["cgst_paise"] for r in ls))),
                      ("SGST", money.to_rupees(sum(r["sgst_paise"] for r in ls))),
                      ("Round Off", money.to_rupees(inv["round_off_paise"])),
                      ("Total", money.to_rupees(t + tx + inv["round_off_paise"]))]
        elif doc_type == "CREDIT_NOTE":
            cn = _get(c, "credit_note", id)
            ls = c.execute("SELECT l.*, i.invoice_no FROM cn_lines l LEFT JOIN invoices i ON i.id = l.invoice_id"
                           " WHERE cn_id = ? ORDER BY line_no", (cn["id"],)).fetchall()
            pairs = [("Document", "CREDIT NOTE"), ("Credit Note No", cn["cn_no"]), ("Credit Note Date", cn["cn_date"]),
                     ("Customer", cn["party_id"]), ("Reason", cn["reason"]), ("IRN", cn["irn"])]
            pairs += [(f"Line {r['line_no']}", f"invoice {r['invoice_no'] or '-'} | taxable "
                                               f"{money.to_rupees(r['taxable_paise'])}") for r in ls]
            pairs.append(("Total", money.to_rupees(sum(r["taxable_paise"] + r["igst_paise"] + r["cgst_paise"]
                                                       + r["sgst_paise"] for r in ls))))
        elif doc_type == "EWAY_BILL":
            e = x.execute("SELECT * FROM ewaybills WHERE ewb_no = ?", (str(id),)).fetchone()
            if e is None:
                inv = _get(c, "invoice", id)
                e = x.execute("SELECT * FROM ewaybills WHERE doc_no = ? AND supplier_gstin = ? ORDER BY generated_at DESC"
                              " LIMIT 1", (inv["invoice_no"], inv["supplier_gstin"])).fetchone()
            if e is None:
                raise LookupError("NOT_FOUND: no e-way bill on the portal for that number or invoice")
            pairs = [("Document", "E-WAY BILL"), ("EWB No", e["ewb_no"]), ("Generated", e["generated_at"]),
                     ("Valid Upto", e["valid_upto"]), ("Doc No", e["doc_no"]), ("Doc Date", e["doc_date"]),
                     ("From GSTIN", e["supplier_gstin"]), ("To GSTIN", e["to_gstin"]), ("From Pincode", e["from_pincode"]),
                     ("To Pincode", e["to_pincode"]), ("Distance km", e["distance_km"]),
                     ("Value", money.to_rupees(e["value_paise"])), ("HSN", e["hsn"]), ("Vehicle No", e["vehicle_no"]),
                     ("Transport Doc No", e["transport_doc_no"]), ("Reason", e["reason"]), ("Status", e["status"])]
        else:
            m = c.execute("SELECT * FROM movements WHERE id = ?", (id,)).fetchone()
            if m is None:
                inv = _get(c, "invoice", id)
                m = c.execute("SELECT m.* FROM movements m JOIN invoice_movements im ON im.movement_id = m.id"
                              " WHERE im.invoice_id = ? ORDER BY m.moved_on LIMIT 1", (inv["id"],)).fetchone()
            if m is None:
                raise LookupError("NOT_FOUND: no movement (LR/POD) for that id or invoice")
            refs = [r[0] for r in c.execute("SELECT i.invoice_no FROM invoice_movements im JOIN invoices i"
                                            " ON i.id = im.invoice_id WHERE im.movement_id = ? ORDER BY i.invoice_no", (m["id"],))]
            ewbs = sorted({r[0] for r in c.execute("SELECT ewb_no FROM invoice_movements WHERE movement_id = ? AND ewb_no"
                                                   " IS NOT NULL", (m["id"],))})
            qty = sum(r[0] for r in c.execute("SELECT qty FROM movement_lines WHERE movement_id = ?", (m["id"],)))
            if doc_type == "LR":
                pairs = [("Document", "LORRY RECEIPT"), ("LR No", m["lr_no"]), ("LR Date", m["lr_date"]),
                         ("Transporter", m["transporter_id"]), ("Vehicle No", m["vehicle_no"]), ("From Pincode", m["from_pincode"]),
                         ("To Pincode", m["to_pincode"]), ("Consignee", m["party_id"]), ("Invoice Refs", ", ".join(refs)),
                         ("EWB No", ", ".join(ewbs)), ("Total Qty", qty)]
            else:
                pairs = [("Document", "PROOF OF DELIVERY"), ("LR No", m["lr_no"]), ("Delivered On", m["delivered_on"]),
                         ("Received By", m["received_by"]), ("Consignee", m["party_id"]), ("Invoice Refs", ", ".join(refs)),
                         ("Total Qty", qty)]
    finally:
        c.close()
        x.close()
    text = _kv(pairs)
    out = {"doc_type": doc_type, "text": text, "sha256": hashlib.sha256(text.encode()).hexdigest()}
    if format == "pdf":
        import fitz
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((50, 60), text, fontsize=9)
        data = pdf.tobytes()
        out.update({"pdf_base64": base64.b64encode(data).decode(), "sha256": hashlib.sha256(data).hexdigest()})
        del out["text"]
    return _ok(out, ("upload_evidence", {"request_id": "REQ-…", "from_erp": True}, "answer an evidence request"))

# ---------------------------------------------------------------- read side shared by client and audit roles

def trace_sale(role, invoice):
    """One sale end to end: enquiry -> SO/pricing/scheme -> movements/LR/POD -> invoice/IRN -> EWB -> GSTR-1 ->
    settlements/bank -> CNs -> journals. Read-only for every role; missing links come back as `gaps`."""
    c = db.connect("erp", role)
    x = db.connect("ext", role)
    try:
        inv = dict(_get(c, "invoice", invoice))
        lines = [dict(r) for r in c.execute("SELECT * FROM invoice_lines WHERE invoice_id = ? ORDER BY line_no", (inv["id"],))]
        so = dict(c.execute("SELECT * FROM sales_orders WHERE id = ?", (inv["so_id"],)).fetchone() or {}) if inv["so_id"] else {}
        so_lines = [dict(r) for r in c.execute("SELECT * FROM so_lines WHERE so_id = ?", (inv["so_id"],))] if so else []
        movs = [dict(r, ewb_no_typed=r["ewb_no"]) for r in c.execute(
            "SELECT m.*, im.ewb_no FROM movements m JOIN invoice_movements im ON im.movement_id = m.id"
            " WHERE im.invoice_id = ? ORDER BY m.moved_on", (inv["id"],))]
        for m in movs:
            m["lines"] = [dict(r) for r in c.execute("SELECT * FROM movement_lines WHERE movement_id = ?", (m["id"],))]
        irn = dict(x.execute("SELECT * FROM einvoices WHERE irn = ?", (inv["irn"],)).fetchone() or {}) if inv["irn"] else {}
        ewbs = [dict(r) for r in x.execute("SELECT * FROM ewaybills WHERE doc_no = ? AND supplier_gstin = ?",
                                           (inv["invoice_no"], inv["supplier_gstin"]))]
        gstr1 = [dict(r) for r in x.execute("SELECT * FROM gstr1_docs WHERE doc_no = ? AND gstin = ?",
                                            (inv["invoice_no"], inv["supplier_gstin"]))]
        sets = [dict(r) for r in c.execute("SELECT * FROM settlements WHERE doc_type = 'INVOICE' AND doc_id = ?", (inv["id"],))]
        bank = [dict(x.execute("SELECT * FROM bank_txns WHERE id = ?", (s["bank_txn_id"],)).fetchone() or {}) for s in sets]
        cns = [dict(r) for r in c.execute("SELECT n.*, l.taxable_paise AS line_taxable FROM cn_lines l JOIN credit_notes n"
                                          " ON n.id = l.cn_id WHERE l.invoice_id = ?", (inv["id"],))]
        jrn = [dict(r) for r in c.execute("SELECT * FROM journals WHERE (source_type = 'invoice' AND source_id = ?) OR"
                                          " (source_type = 'credit_note' AND source_id IN (SELECT cn_id FROM cn_lines"
                                          " WHERE invoice_id = ?)) ORDER BY no", (inv["id"], inv["id"]))]
        for j in jrn:
            j["lines"] = [dict(r) for r in c.execute("SELECT * FROM journal_lines WHERE journal_no = ? ORDER BY line_no",
                                                     (j["no"],))]
    finally:
        c.close()
        x.close()
    t = sum(l["taxable_paise"] for l in lines)
    total = t + sum(l["igst_paise"] + l["cgst_paise"] + l["sgst_paise"] for l in lines)
    goods = any(not _is_service(l["hsn"]) and not l["item_id"].startswith("CHG") for l in lines)
    gaps = []
    if not so:
        gaps.append("NO_SALES_ORDER")
    if goods and not movs:
        gaps.append("NO_MOVEMENT")
    if goods and so.get("delivery_terms") == "FOR_DESTINATION" and not any(m["delivered_on"] for m in movs):
        gaps.append("NO_POD")
    if inv["buyer_gstin"] and not inv["irn"]:
        gaps.append("NO_IRN")
    if inv["irn"] and not irn:
        gaps.append("IRN_NOT_ON_PORTAL")
    if goods and total > config.EWB_THRESHOLD_PAISE and not any(e["status"] == "ACT" for e in ewbs):
        gaps.append("NO_ACTIVE_EWB")
    if not gstr1:
        gaps.append("NOT_IN_GSTR1")
    if not sets:
        gaps.append("UNPAID")
    if not jrn:
        gaps.append("NO_JOURNAL")
    return {"invoice": inv, "lines": lines, "sales_order": so, "so_lines": so_lines, "movements": movs,
            "einvoice": irn, "ewaybills": ewbs, "gstr1": gstr1, "settlements": sets, "bank_txns": bank,
            "credit_notes": cns, "journals": jrn, "taxable_paise": t, "total_paise": total + inv["round_off_paise"],
            "gaps": gaps}
