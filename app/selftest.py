"""End-to-end self-check: python -m app.selftest [--keep] (07 §8.6).

Runs in a throwaway data dir, so it never touches data/. Today it proves the platform (schema, role-bound handles,
append-only guards, clock), checks every 07 §4 contract signature, smoke-tests the pure cores that already exist,
and lists the 07 §8.6 end-to-end steps, marked SKIP while the module that owns each one is still a STUB.
SH-6 fills the E2E steps in as the queue lands. Exit code 1 if anything FAILs.
"""
import importlib
import inspect
import json
import random
import shutil
import sqlite3
import sys
import tempfile
import traceback

from app import config

RESULTS = []
SELFTEST_IDS = (("adm.kiran", "client_admin"), ("aud.judge", "audit_staff"), ("ptn.mehta", "engagement_partner"))


def step(name):
    def deco(fn):
        try:
            out = fn()
            RESULTS.append(("SKIP" if isinstance(out, str) and out.startswith("SKIP") else "PASS", name,
                            out if isinstance(out, str) else ""))
        except Exception as e:  # noqa: BLE001 - a selftest reports every failure and carries on
            RESULTS.append(("FAIL", name, f"{type(e).__name__}: {e}"))
            if "-v" in sys.argv:
                traceback.print_exc()
        return fn
    return deco


def expect_error(fn, *fragments):
    try:
        fn()
    except Exception as e:  # noqa: BLE001
        msg = str(e)
        assert all(f in msg for f in fragments), f"wrong error: {msg!r}"
        return
    raise AssertionError(f"expected an error containing {fragments}")

# ---------------------------------------------------------------- contracts (07 §4)

ACTOR = ("actor", "role")
CONTRACTS = {
    "db": {"connect": ("db", "role"), "init_all": ("reset",), "tx": ("conn",)},
    "chain": {"append": ("conn", "table", "row"), "verify": ("conn", "table", "anchors")},
    "money": {"to_paise": ("s",), "fmt_inr": ("paise",), "tax_split": ("taxable_paise", "rate_bp", "inter_state")},
    "ledger": {"post_journal": ("conn", "voucher_date", "voucher_type", "lines", "narration", "actor", "role", "tool",
                                "source_type", "source_id", "reverses_no", "aje_ref"),
               "reverse_journal": ("conn", "journal_no", "date", "actor", "role", "tool", "reason"),
               "trial_balance": ("conn", "upto"), "seed_coa": ("conn",)},
    "erp": {"write_doc": ("conn", "actor", "role", "tool", "action", "doc_type", "doc_id", "reason"),
            **{s: ACTOR for s in ("create_party", "link_parties", "create_item", "create_lead", "create_scheme",
                                  "create_sales_order", "issue_invoice", "generate_eway_bill", "record_dispatch",
                                  "record_delivery", "issue_credit_note", "record_receipt", "record_payment",
                                  "record_purchase", "post_journal_entry", "file_gstr1", "file_gstr3b", "amend_record",
                                  "cancel_document", "import_records", "export_sales_register", "render_document")}},
    "portal": {"register_irn": ("supplier_gstin", "fy", "doc_type", "doc_no", "doc_date", "buyer_gstin",
                                "taxable_paise", "tax_paise"),
               "cancel_irn": ("irn",), "generate_eway_bill": (), "cancel_eway_bill": ("ewb_no",),
               "file_gstr1": ("gstin", "period", "omit", "changes"), "file_gstr3b": ("gstin", "period", "figures"),
               "bank_txn": ("account_no", "value_date", "amount_paise", "cp_name", "cp_account", "cp_ifsc", "utr",
                            "narration"),
               "verify_signature": ("einvoice_row",)},
    "trail": {"trail": ("principal", "role", "tool", "args", "outcome", "result", "transport"),
              "anchor": ("engagement_id", "label")},
    "notify": {"emit": ("kind", "payload", "engagement_id")},
    "intake": {"submit": ("kind", "content", "filename", "submitted_by", "request_id", "declared"),
               "on_client_upload": ("submission_id",)},
    "register": {"parse_verbatim": ("content", "filename"), "map_columns": ("headers",),
                 "reconcile": ("engagement_id",)},
    "related_parties": {"run_d1_d2": ("engagement_id",), "rp_status": ("party_id", "on_date")},
    "analytics": {"run_all": ("engagement_id", "detectors", "incremental")},
    "sampling": {"select": ("engagement_id", "kind", "size", "seed", "actor"), "reproduce": ("run_id",),
                 "pps": ("items", "k", "rng")},
    "risk_graph": {"build": ("engagement_id",)},
    "findings": {"from_hits": ("engagement_id",), "aje": ("finding",), "sa450_summary": ("engagement_id",)},
    "evidence": {"extract": ("submission_id",), "test_sample": ("run_id", "sample_no", "actor", "note"),
                 "bundle": ("invoice_id",), "check_for_changes": ("engagement_id",),
                 "verify_integrity": ("engagement_id",)},
    "workpaper": {"build": ("engagement_id",)},
    "memo": {"build": ("engagement_id",)},
}
OTHER_MODULES = ("config", "playbook", "gateway", "tools_client", "tools_audit", "mcp_server", "investigator", "seed")


def mod(name):
    return importlib.import_module(f"app.{name}")


def is_stub(name):
    return getattr(mod(name), "STUB", False)


def _params(fn):
    ps = [p for p in inspect.signature(fn).parameters.values() if p.kind not in (p.VAR_KEYWORD, p.VAR_POSITIONAL)]
    return ps


def check_signature(module, fname, want):
    fn = getattr(mod(module), fname, None)
    assert callable(fn), f"{module}.{fname} missing"
    ps = _params(inspect.unwrap(fn))
    got = tuple(p.name for p in ps[:len(want)])
    assert got == want, f"{module}.{fname}{got} != contract {want}"
    extra = [p.name for p in ps[len(want):] if p.default is inspect.Parameter.empty]
    assert not extra, f"{module}.{fname} has extra required params {extra}"

# ---------------------------------------------------------------- steps

def run():
    from app import db

    @step("imports: every app module loads")
    def _():
        for m in list(CONTRACTS) + list(OTHER_MODULES):
            mod(m)
        stubs = sorted(m for m in list(CONTRACTS) + list(OTHER_MODULES) if is_stub(m))
        return f"stubs: {', '.join(stubs)}" if stubs else ""

    @step("contracts: 07 §4 signatures")
    def _():
        n = 0
        for module, fns in CONTRACTS.items():
            for fname, want in fns.items():
                check_signature(module, fname, want)
                n += 1
        return f"{n} functions"

    @step("SH-2 init_all(reset=True): 4 databases, every 03 table, WAL")
    def _():
        db.init_all(reset=True)
        want = {"erp": 20, "ext": 6, "pbc": 3, "audit": 13}
        for name, n in want.items():
            c = sqlite3.connect(str(config.DB_PATHS[name]))
            tables = c.execute("SELECT count(*) FROM sqlite_master WHERE type='table'").fetchone()[0]
            mode = c.execute("PRAGMA journal_mode").fetchone()[0]
            c.close()
            assert tables == n, f"{name}.db has {tables} tables, 03 lists {n}"
            assert mode == "wal", f"{name}.db journal_mode={mode}"
        db.init_all(reset=False)                               # idempotent re-run
        return "erp 20 · ext 6 · pbc 3 · audit 13"

    @step("SH-2 independence: audit role cannot write erp/ext (readonly database)")
    def _():
        for name, sql in (("erp", "INSERT INTO accounts(code,name,type) VALUES ('X','x','ASSET')"),
                          ("ext", "INSERT INTO bank_txns(id,account_no,value_date,amount_paise) VALUES ('B','1','2025-04-01',1)")):
            for role in ("audit_staff", "engagement_partner", "system"):
                c = db.connect(name, role)
                expect_error(lambda: c.execute(sql), "readonly database")
                c.close()

    @step("SH-2 independence: client roles have no audit.db handle; orchestrator none on erp.db")
    def _():
        for role in ("sales", "warehouse", "accounts", "client_admin", "portal"):
            expect_error(lambda: db.connect("audit", role), "PERMISSION_DENIED")
        expect_error(lambda: db.connect("erp", "orchestrator"), "PERMISSION_DENIED")
        expect_error(lambda: db.connect("erp", "hacker"), "PERMISSION_DENIED")
        c = db.connect("audit", "orchestrator")
        expect_error(lambda: c.execute("INSERT INTO events(kind,at) VALUES ('SIGNED_OFF','x')"), "readonly")
        c.close()

    @step("SH-2 writers: client writes erp, portal writes ext, tx() rolls back on error")
    def _():
        c = db.connect("erp", "client_admin")
        try:
            with db.tx(c):
                c.execute("INSERT INTO accounts(code,name,type) VALUES ('ZZ_SELFTEST','x','ASSET')")
                raise RuntimeError("rollback")
        except RuntimeError:
            pass
        assert c.execute("SELECT count(*) FROM accounts WHERE code='ZZ_SELFTEST'").fetchone()[0] == 0
        c.close()
        p = db.connect("ext", "portal")
        with db.tx(p):
            p.execute("INSERT INTO bank_txns(id,account_no,value_date,amount_paise) VALUES ('B-ST','1','2025-04-01',-5)")
        p.close()
        c = db.connect("ext", "client_admin")
        expect_error(lambda: c.execute("DELETE FROM bank_txns"), "readonly database")
        c.close()

    @step("SH-2 guards: append-only + [F] triggers abort")
    def _():
        c = db.connect("erp", "client_admin")
        with db.tx(c):
            c.execute("INSERT INTO erp_log(recorded_at,actor,role,tool,action,doc_type,doc_id,after_json,prev_hash,hash)"
                      " VALUES ('t','a','r','t','CREATE','x','1','{}','p0','h1')")
            c.execute("INSERT INTO parties(id,kind,name,recorded_by,recorded_at) VALUES ('P-ST','ENTITY','x','a','t')")
            c.execute("INSERT INTO invoices(id,invoice_no,series,supplier_gstin,invoice_date,party_id,place_of_supply)"
                      " VALUES ('I-ST','G2/1/25','G2','g','2025-04-01','P-ST','27')")
        expect_error(lambda: c.execute("UPDATE erp_log SET actor='b'"), "append-only: erp_log")
        expect_error(lambda: c.execute("DELETE FROM erp_log"), "append-only: erp_log")
        expect_error(lambda: c.execute("UPDATE invoices SET invoice_date='2025-04-02'"), "financial document")
        expect_error(lambda: c.execute("DELETE FROM invoices"), "never deleted")
        c.execute("UPDATE invoices SET status='CANCELLED', cancelled_on='2025-04-02'")   # cancel is allowed
        c.close()
        a = db.connect("audit", "audit_staff")
        a.execute("INSERT INTO audit_trail(at,principal,role,transport,tool,outcome,prev_hash,hash)"
                  " VALUES ('t','p','audit_staff','SYSTEM','whoami','OK','p0','h1')")
        expect_error(lambda: a.execute("UPDATE audit_trail SET tool='x'"), "append-only: audit_trail")
        expect_error(lambda: a.execute("DELETE FROM audit_trail"), "append-only: audit_trail")
        a.close()
        db.init_all(reset=True)                                    # leave a clean world for the steps below

    @step("config: settable clock in IST, unit master, risk weights")
    def _():
        config.set_clock("2026-03-31T23:30:00+05:30")
        assert config.now() == "2026-03-31T23:30:00+05:30" and config.today() == "2026-03-31"
        config.set_clock("2026-03-31T20:00:00+00:00")             # 01:30 IST on 1 Apr is April (05 D3, C7)
        assert config.today() == "2026-04-01"
        config.set_clock(None)
        assert config.now().endswith("+05:30")
        from app.related_parties import gstin_ok
        assert all(gstin_ok(u["gstin"]) and u["gstin"][2:12] == config.SELF_PAN for u in config.UNITS.values())
        assert config.risk_weight("FORMAT_ONLY") == 0.2 and config.risk_weight("NO_MOVEMENT") == 1.0

    @step("pure cores: seed plan -> D1/D2 analyse, register parse + D0")
    def _():
        from app import register, related_parties, seed
        plan = seed.build_plan(50, 7)
        res = related_parties.analyse(plan["parties"], plan["links"], [], config.PERIOD_START, config.PERIOD_END)
        und = [h for h in res["hits"] if h["code"] == "UNDISCLOSED_RP"]
        assert und, "undeclared group companies should surface when the RP list is empty"
        headers, rows = register.parse_verbatim(b"Bill no.,Date,Total Cost\nG2/1/25,01-04-2025,100\n", "r.csv")
        hits, _p = register.diff_register([register.map_row(headers, r) for r in rows], [], *(
            config.PERIOD_START, config.PERIOD_END))
        assert [h["code"] for h in hits] == ["NOT_IN_SOR"]
        return f"{len(plan['parties'])} parties, {len(und)} undisclosed with an empty RP list"

    @step("stubs return their contract shapes")
    def _():
        checks = {
            "analytics": lambda m: set(m.run_all("E")) >= {"counts", "top"},
            "sampling": lambda m: set(m.select("E", "sales_tod")) >= {"run_id", "seed", "items"}
                                  and isinstance(m.pps([("I1", 1)], 1, random.Random(1)), list),
            "evidence": lambda m: set(m.verify_integrity("E")) >= {"ok", "checks"}
                                  and set(m.check_for_changes("E")) >= {"changes", "stale_samples"},
            "risk_graph": lambda m: set(m.build("E")) >= {"nodes", "edges", "cycles", "company"},
            "workpaper": lambda m: set(m.build("E")) >= {"path", "sha256", "download_key"},
            "memo": lambda m: set(m.build("E")) >= {"path", "sha256", "download_key"},
            "findings": lambda m: isinstance(m.from_hits("E"), list),
        }
        done = []
        for name, ok in checks.items():
            if is_stub(name):
                assert ok(mod(name)), f"{name} stub shape"
                done.append(name)
        return f"checked: {', '.join(done)}" if done else "no stubs left"

    @step("ERP: one sale end to end (SO -> invoice/IRN -> EWB -> dispatch -> POD -> GSTR-1 -> receipt), trace has no gaps")
    def _():
        from app import erp
        if is_stub("erp"):
            return "SKIP (stub: erp)"
        A, W, C = ("acc.neha", "accounts"), ("whs.ganesh", "warehouse"), ("adm.kiran", "client_admin")
        config.set_clock("2026-03-01T10:00:00+05:30")
        erp.create_party(*C, id="SELF", tags="SELF", name=config.ENTITY_NAME, gstin=config.UNITS["G2"]["gstin"],
                         bank_account="50100012345678", ifsc="HDFC0001234")
        erp.create_party(*C, id="P-0001", name="Selftest Customer Pvt Ltd", gstin="27AAPFU0939F1ZV", pincode="400001")
        erp.create_item(*C, id="IT-001", name="Valve", hsn="8481", list_rate_paise=680000)
        so = erp.create_sales_order("sales.priya", "sales", party_id="P-0001", confirmed_on="2026-03-05",
                                    lines=[{"item_id": "IT-001", "qty": 10}])["result"]["id"]
        inv = erp.issue_invoice(*A, so_id=so, invoice_date="2026-03-28")["result"]
        ewb = erp.generate_eway_bill(*W, invoice_id=inv["id"], vehicle_no="MH08AB1234", distance_km=300)["result"]
        mv = erp.record_dispatch(*W, invoice_nos=[inv["invoice_no"]], moved_on="2026-03-28", lr_no="LR-1",
                                 ewb_no=ewb["ewb_no"])["result"]["id"]
        erp.record_delivery(*W, movement_id=mv, delivered_on="2026-03-29")
        config.set_clock("2026-04-11T10:00:00+05:30")
        erp.file_gstr1(*A, period="202603")
        erp.record_receipt(*A, amount_paise=inv["total_paise"], invoice_nos=[inv["id"]], settled_on="2026-04-20")
        config.set_clock(None)
        gaps = erp.trace_sale("audit_staff", inv["invoice_no"])["gaps"]
        assert gaps == [], gaps
        db.init_all(reset=True)
        return inv["invoice_no"]

    # 07 §8.6 end to end, in order; each step needs the listed modules to be real
    W = {}                                       # state shared by the E2E steps

    def e2e_seed():
        import time
        from app import seed
        t0 = time.perf_counter()
        W["plan"] = seed.build_plan(300, 7)
        seed.load(W["plan"], reset=True)
        config.set_clock("2026-10-05T10:00:00+05:30")
        from app import chain, erp
        with db.borrow("erp", "audit_staff") as e:
            assert chain.verify(e, "erp_log")["valid"] and erp.out_of_band(e) == []
            n = e.execute("SELECT count(*) FROM invoices").fetchone()[0]
        return f"{n} invoices in {time.perf_counter() - t0:.1f}s"

    def e2e_upload():
        from app import erp, intake
        exp = erp.export_sales_register("acc.neha", "accounts")["result"]
        content, name = intake.build_register(exp["export_id"], edits=[{"row_no": 2, "column": "Total Cost", "value": 1}],
                                              drop_rows=[5], add_rows=[{"Bill no.": "G2/9999/25", "Date": "15-01-2026",
                                                                        "Customer": "Ghost Traders", "Total Cost": 5000}])
        W["upload"] = intake.on_client_upload(intake.submit("SALES_REGISTER", content, name, "acc.neha"))
        return f"{W['upload']['rows']} rows, {W['upload']['d0_hits']} D0 hits"

    def e2e_d0():
        with db.borrow("audit", "audit_staff") as a:
            codes = {r[0] for r in a.execute("SELECT code FROM risk_hits WHERE detector = 'D0'")}
            rp = a.execute("SELECT count(*) FROM risk_hits WHERE detector IN ('D1','D2')").fetchone()[0]
        assert {"NOT_IN_SOR", "MISSING_FROM_REGISTER", "TAXABLE_DIFF"} <= codes, codes
        return f"D0 codes {sorted(codes)}; {rp} D1/D2 hits (no RP list uploaded: group companies undisclosed)"

    def e2e_mcp():
        import asyncio
        import hashlib
        from app import gateway, mcp_server  # noqa: F401
        gateway.STDIO_ID = ("aud.judge", "audit_staff")
        seen = {t.name for t in asyncio.run(gateway.mcp.list_tools())}
        assert "get_engagement_status" in seen and "issue_invoice" not in seen
        erp_file = config.DB_PATHS["erp"]
        before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (erp_file, erp_file.with_name("erp.db-wal"))
                  if p.exists()]
        try:
            asyncio.run(gateway.mcp.call_tool("create_party", {"name": "Ghost"}))
            raise AssertionError("an auditor created a party")
        except Exception as e:  # noqa: BLE001
            assert "PERMISSION_DENIED" in str(e), e
        after = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (erp_file, erp_file.with_name("erp.db-wal"))
                 if p.exists()]
        assert before == after
        st = json.loads(asyncio.run(gateway.mcp.call_tool("get_engagement_status", {})).content[0].text)
        gateway.STDIO_ID = None
        with db.borrow("audit", "audit_staff") as a:
            outcomes = [r[0] for r in a.execute("SELECT outcome FROM audit_trail ORDER BY seq")]
        assert outcomes[-2:] == ["DENIED", "OK"], outcomes[-2:]
        return f"{len(seen)} audit tools; denied call trailed, erp.db byte-identical; status: {st['summary']['status']}"

    def e2e_analytics():
        from app import analytics
        r = analytics.run_all(config.ENGAGEMENT_ID)
        W["counts"] = r["counts"]
        serious = sum(n for d, codes in r["counts"].items() for c, n in codes.items()
                      if d not in ("D0", "D1", "D2", "D5", "D10"))
        d10 = sum(r["counts"].get("D10", {}).values())
        return (f"detectors {sorted(r['counts'])}; {serious} D3-D14 hits on the clean year; {d10} D10 hits = the seed's "
                f"planted trading network")

    def e2e_samples():
        from app import sampling
        out = []
        for kind in ("sales_tod", "cutoff", "credit_note"):
            try:
                r = sampling.select(config.ENGAGEMENT_ID, kind, actor="aud.judge")
            except ValueError as e:
                if kind == "credit_note" and "empty" in str(e):
                    out.append("CN: none in a clean year")
                    continue
                raise
            assert sampling.reproduce(r["run_id"])["identical"], kind
            W.setdefault("runs", {})[kind] = r
            out.append(f"{kind} {r['size']} ({r['requests_sent']} requests)")
        return "; ".join(out) + "; reproduce identical"

    def e2e_evidence():
        from app import erp, evidence, intake
        run = W["runs"]["sales_tod"]
        with db.borrow("erp", "audit_staff") as e, db.borrow("ext", "audit_staff") as x:
            nos = {r[0]: r[1] for r in e.execute("SELECT id, invoice_no FROM invoices")}
            pick = next(i for i in run["items"] if x.execute("SELECT 1 FROM ewaybills WHERE doc_no = ?",
                                                             (nos[i["doc_id"]],)).fetchone())
        no = nos[pick["doc_id"]]
        with db.borrow("pbc", "accounts") as p:
            reqs = {r[1]: r[0] for r in p.execute("SELECT id, doc_kind FROM evidence_requests WHERE about_doc_no = ?", (no,))}
        for kind, rid in reqs.items():
            text = erp.render_document("acc.neha", "accounts", doc_type=kind, id=no)["result"]["text"]
            if kind == "EWAY_BILL":                   # the client "improves" the e-way bill copy
                text = "\n".join("Value: 1.00" if ln.startswith("Value:") else ln for ln in text.splitlines())
            intake.on_client_upload(intake.submit("EVIDENCE", text.encode(), f"{kind}.txt", "acc.neha", request_id=rid))
        r = evidence.test_sample(run["run_id"], pick["sample_no"], "aud.judge")
        assert "DOC_FIELD_MISMATCH" in {x["code"] for x in r["exceptions"]}, r["exceptions"]
        W["tested"] = (run["run_id"], pick["sample_no"], pick["doc_id"])
        return f"sample {pick['sample_no']} ({no}): forged EWB copy -> DOC_FIELD_MISMATCH; conclusion {r['conclusion_proposed']}"

    def e2e_stale():
        from app import erp, evidence
        run_id, n, inv = W["tested"]
        with db.borrow("erp", "audit_staff") as e:
            mv = e.execute("SELECT movement_id FROM invoice_movements WHERE invoice_id = ?", (inv,)).fetchone()[0]
        erp.amend_record("whs.ganesh", "warehouse", entity="movements", id=mv, changes={"lr_date": "2026-04-02"},
                         reason="LR date corrected")
        ch = evidence.check_for_changes(config.ENGAGEMENT_ID)
        item = next(c for c in ch["changes"] if c["kind"] == "EVIDENCE_FIELD_CHANGED" and c["pk"] == mv)
        assert f"{run_id}#{n}" in item["affected_samples"] and item["by"] == "whs.ganesh"
        return f"{run_id}#{n} STALE: lr_date {item['fields']['header.lr_date']} by {item['by']} via {item['tool']}"

    def e2e_deliverables():
        from app import findings, memo, workpaper
        run_id, n, _inv = W["tested"]
        from app import evidence
        first = evidence.test_sample(run_id, n, "aud.judge")                 # sees the LR change: STALE
        again = evidence.test_sample(run_id, n, "aud.judge", note="re-tested after the LR date change")
        assert first["conclusion_proposed"] == "STALE" and again["conclusion_proposed"] != "STALE"
        new = findings.from_hits(config.ENGAGEMENT_ID)
        with db.borrow("audit", "audit_staff") as a:
            ajes = [json.loads(r[0]) for r in a.execute("SELECT proposed_aje_json FROM findings WHERE proposed_aje_json"
                                                        " IS NOT NULL")]
        assert all(sum(x["dr_paise"] for x in j) == sum(x["cr_paise"] for x in j) for j in ajes)
        wp, mm = workpaper.build(config.ENGAGEMENT_ID), memo.build(config.ENGAGEMENT_ID)
        from openpyxl import load_workbook
        assert load_workbook(wp["path"]).sheetnames == ["Sales Reg.", "Sales TOD", "Sales Cut-off TOD", "Credit Note TOD"]
        return f"{len(new)} PROPOSED findings ({len(ajes)} with balanced AJEs); workpaper + memo (DRAFT) built"

    def e2e_signoff():
        import subprocess
        from app import findings, trail
        partner = {"principal": "ptn.mehta", "role": "engagement_partner", "name": "CA Rakesh Mehta",
                   "membership_no": "118204"}
        with db.borrow("audit", "audit_staff") as a:
            ids = [r[0] for r in a.execute("SELECT id FROM findings WHERE status = 'PROPOSED' ORDER BY id")]
        for fid in ids:
            findings.decide(fid, "CONFIRMED", "agreed with the evidence and the reasons", partner)
        r = findings.sign_off(config.ENGAGEMENT_ID, partner)
        out = config.AUDIT_FILES_DIR / config.ENGAGEMENT_ID
        trail.export(out, config.ENGAGEMENT_ID)
        v = subprocess.run([sys.executable, str(config.ROOT / "scripts" / "verify_trail.py"), str(out / "audit_trail.jsonl"),
                            str(out / "erp_log_head.json")], capture_output=True, text=True)
        assert v.returncode == 0, v.stdout + v.stderr
        return f"{len(ids)} findings confirmed; SIGNED_OFF {r['signoff_hash'][:12]}; verify_trail: {v.stdout.strip()}"

    def e2e_integrity():
        import sqlite3
        from app import evidence, trail
        trail.anchor(config.ENGAGEMENT_ID, "PERIODIC")
        assert evidence.verify_integrity(config.ENGAGEMENT_ID)["ok"]
        raw = sqlite3.connect(str(config.DB_PATHS["erp"]), isolation_level=None)
        raw.execute("DROP TRIGGER invoices_f")
        raw.execute("UPDATE invoices SET invoice_date = '2026-04-02' WHERE id = 'INV-000002'")
        raw.close()
        v = evidence.verify_integrity(config.ENGAGEMENT_ID)
        assert not v["ok"] and "OUT_OF_BAND_EDIT" in {p["code"] for p in v["problems"]}
        return "clean -> OK; direct SQLite edit -> OUT_OF_BAND_EDIT"

    def e2e_risk():
        from app import risk_graph
        g = risk_graph.build(config.ENGAGEMENT_ID)
        flagged = sorted({c for cyc in g["cycles"] if cyc["flagged"] for c in cyc["codes"]})
        assert {"ROUND_TRIP_GOODS", "ROUND_TRIP_CASH"} <= set(flagged), flagged
        assert any(c["codes"] == ["RECIPROCAL_EXPLAINED"] and not c["flagged"] for c in g["cycles"])
        co = g["company"]
        return (f"company {co['band']} (families {','.join(co['families']) or '-'}, circular "
                f"{co['circular_paise'] // 100:,}); cycles flagged {flagged}; job-work pair explained")

    e2e = [("seed 300 invoices through erp services", ("erp", "ledger", "portal", "chain"), e2e_seed),
           ("export register (Fig. 1) + doctored upload", ("erp", "intake"), e2e_upload),
           ("D0 reconcile + D1/D2 on the frozen register", ("intake",), e2e_d0),
           ("MCP: auditor sees only audit tools; a client tool is denied + trailed; erp.db unchanged",
            ("gateway", "trail"), e2e_mcp),
           ("run_analytics", ("analytics",), e2e_analytics),
           ("select_samples x3 + reproduce identical", ("sampling", "trail"), e2e_samples),
           ("render/upload evidence (one forged) -> test_sample flags it", ("evidence",), e2e_evidence),
           ("amend LR date after testing -> check_for_changes names the sample", ("evidence", "erp"), e2e_stale),
           ("risk graph: cycles by minimum-span decomposition, company band", ("risk_graph",), e2e_risk),
           ("findings + AJEs, workpaper xlsx, memo", ("findings", "workpaper", "memo"), e2e_deliverables),
           ("partner confirm + sign-off; verify_trail OK", ("trail", "findings"), e2e_signoff),
           ("verify_integrity clean, then out-of-band edit detected", ("evidence", "chain"), e2e_integrity)]
    for name, owners, fn in e2e:
        stubs = [o for o in owners if is_stub(o)]

        @step(f"E2E: {name}")
        def _(stubs=stubs, fn=fn):
            if stubs:
                return f"SKIP (stub: {', '.join(stubs)})"
            return fn() if fn else "SKIP (SH-6 wires this step)"


def main():
    import json
    import os
    keep = "--keep" in sys.argv
    tmp = tempfile.mkdtemp(prefix="auditpilot-selftest-")
    config.set_data_dir(tmp)
    ids = os.path.join(tmp, "identities.json")             # the Gateway reads IDENTITIES at import (04 §2)
    with open(ids, "w") as f:
        json.dump({f"{r[:3]}_selftest": {"principal": p, "role": r, "name": p} for p, r in SELFTEST_IDS}, f)
    os.environ["IDENTITIES"] = ids
    try:
        run()
    finally:
        config.set_clock(None)
        if not keep:
            shutil.rmtree(tmp, ignore_errors=True)
    width = max(len(n) for _, n, _ in RESULTS)
    for status, name, note in RESULTS:
        print(f"{status:<4}  {name:<{width}}  {note}")
    counts = {s: sum(1 for r in RESULTS if r[0] == s) for s in ("PASS", "FAIL", "SKIP")}
    print(f"\nselftest: {counts['PASS']} passed, {counts['FAIL']} failed, {counts['SKIP']} skipped"
          f"{f'  (data kept in {tmp})' if keep else ''}")
    return 1 if counts["FAIL"] else 0


if __name__ == "__main__":
    sys.exit(main())
