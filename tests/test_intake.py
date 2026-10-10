"""DA-1 load + DA-2 intake/reconcile + DA-3 run_d1_d2 on a real (small) seeded world."""
import json
from collections import Counter

import pytest

from app import chain, config, db, erp, intake, related_parties, seed

ACC, ADM = ("acc.neha", "accounts"), ("adm.kiran", "client_admin")


@pytest.fixture()
def seeded(world):
    plan = seed.build_plan(40, 7)
    seed.load(plan, reset=True)
    config.set_clock("2026-10-05T10:00:00+05:30")          # audit fieldwork, months after the year end
    return plan


def upload(kind, content, filename, who=ACC):
    sid = intake.submit(kind, content, filename, who[0])
    return intake.on_client_upload(sid)


def hits(codes_of=None):
    with db.borrow("audit", "audit_staff") as a:
        rows = [dict(r) for r in a.execute("SELECT * FROM risk_hits")]
    return [h for h in rows if not codes_of or h["detector"] in codes_of]


def test_seeded_world_is_consistent(seeded):
    with db.borrow("erp", "audit_staff") as e:
        assert chain.verify(e, "erp_log")["valid"] and erp.out_of_band(e) == []
        assert e.execute("SELECT count(*) FROM invoices").fetchone()[0] == len(seeded["sales"])
        unpaid = e.execute("SELECT count(*) FROM invoices i WHERE NOT EXISTS (SELECT 1 FROM settlements s"
                           " WHERE s.doc_id = i.id)").fetchone()[0]
        assert unpaid == 0
        bal = e.execute("SELECT SUM(dr_paise) - SUM(cr_paise) FROM journal_lines").fetchone()[0]
        assert bal == 0
        assert e.execute("SELECT count(*) FROM purchase_bills").fetchone()[0] == len(seeded["purchases"])
    sample = next(s for s in seeded["sales"] if s["dispatch"] and s["invoice_date"] <= "2026-03-31")
    with db.borrow("erp", "audit_staff") as e:
        inv = e.execute("SELECT id FROM invoices WHERE invoice_date = ? AND party_id = ?",
                        (sample["invoice_date"], sample["party_id"])).fetchone()[0]
    gaps = erp.trace_sale("audit_staff", inv)["gaps"]
    assert gaps == [], gaps


def test_clean_register_upload_has_no_d0_hits(seeded):
    exp = erp.export_sales_register(*ACC)["result"]
    content, name = intake.build_register(exp["export_id"])
    out = upload("SALES_REGISTER", content, name)
    assert out["d0_hits"] == 0 and out["rows"] == exp["rows"] and "doc_no" in out["columns_mapped"]
    with db.borrow("audit", "audit_staff") as a:
        assert a.execute("SELECT count(*) FROM register_rows").fetchone()[0] == exp["rows"]
        assert [r[0] for r in a.execute("SELECT kind FROM events")] == ["REGISTER_UPLOADED"]
        assert a.execute("SELECT label FROM anchors").fetchone()[0] == "UPLOAD"
        eng = a.execute("SELECT * FROM engagements").fetchone()
        assert eng["id"] == config.ENGAGEMENT_ID and eng["register_submission_id"] == out["submission_id"]


def test_doctored_register_every_k_plant(seeded):
    exp = erp.export_sales_register(*ACC)["result"]
    with db.borrow("erp", "audit_staff") as e:
        first = e.execute("SELECT id FROM invoices WHERE invoice_date <= '2026-03-31' ORDER BY invoice_date, invoice_no"
                          " LIMIT 1").fetchone()[0]
    erp.cancel_document(*ACC, doc_type="invoice", id=first, reason="duplicate billing")    # K7: export predates this
    import openpyxl
    with db.borrow("pbc", "audit_staff") as p:
        path = p.execute("SELECT stored_path FROM exports WHERE id = ?", (exp["export_id"],)).fetchone()[0]
    ws = openpyxl.load_workbook(path).active
    row3 = {ws.cell(row=1, column=c).value: ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)}
    other = next(p for p in seeded["parties"] if p["tags"] == "CUSTOMER" and p["gstin"] and p["id"] != "P-0001")
    content, name = intake.build_register(exp["export_id"], edits=[
        {"row_no": 2, "column": "Total Cost", "value": 1},                                   # K3 amount
        {"row_no": 3, "column": "Date", "value": "02-04-2026"},                              # K4 date (and FY)
        {"row_no": 4, "column": "Customer", "value": other["name"]},                         # K5 customer
        {"row_no": 4, "column": "GST No.", "value": other["gstin"]}],
        drop_rows=[5],                                                                       # K2 omitted
        add_rows=[{"Bill no.": "G2/9999/25", "Date": "15-01-2026", "Customer": "Ghost Traders",  # K1 fabricated
                   "Total Cost": 500000, "Gross Amt.": 590000},
                  row3])                                                                     # K6 duplicate of row 3
    upload("SALES_REGISTER", content, name)
    c = Counter(h["code"] for h in hits({"D0"}))
    for code in ("NOT_IN_SOR", "MISSING_FROM_REGISTER", "TAXABLE_DIFF", "DATE_DIFF", "FY_DIFF", "CUSTOMER_DIFF",
                 "GSTIN_DIFF", "DUPLICATE_ROW", "CANCELLED_BUT_LISTED"):
        assert c[code] >= 1, (code, c)
    assert any(h["code"] == "TIEOUT_RESIDUE" and h["object_id"] == "ledger" for h in hits({"D0"}))  # late cancel


def test_rp_list_with_an_omission_and_a_wrong_date(seeded):
    rows = ["Name,PAN,Relation,Effective From"]
    by = {p["id"]: p for p in seeded["parties"]}
    for pid in ("P-DIR1", "P-DIR2", "P-DIR3", "P-DIR4", "P-KMP1", "P-KMP2", "P-REL1", "P-REL2", "P-REL3", "P-REL4",
                "P-GRP2"):
        rows.append(f"{by[pid]['name']},{by[pid]['pan']},declared,01-04-2015")
    rows.append(f"{by['P-GRP1']['name']},{by['P-GRP1']['pan']},director's company,01-01-2026")     # wrong date
    out = upload("RP_LIST", "\n".join(rows).encode(), "rp_list.csv", ADM)                     # P-GRP3 omitted
    assert out["rows"] == 12
    d = {(h["code"], h["object_id"]) for h in hits({"D1", "D2"})}
    assert ("UNDISCLOSED_RP", "P-GRP3") in d and ("RP_PERIOD_MISSTATED", "P-GRP1") in d
    assert not any(c == "UNDISCLOSED_RP" and o.startswith("P-NET") for c, o in d)               # network is not an RP
    assert related_parties.rp_status("P-GRP2", "2025-06-01")[0] == "DECLARED_RP"
    assert related_parties.rp_status("P-GRP1", "2025-06-01")[0] == "UNDISCLOSED_RP"           # in the misstated gap
    assert related_parties.rp_status("P-0001", "2025-06-01")[0] == "NRP"
    with db.borrow("audit", "audit_staff") as a:
        assert json.loads(a.execute("SELECT mapped_json FROM rp_list_rows WHERE row_no = 12").fetchone()[0])[
            "effective_from"] == "2026-01-01"


def test_submissions_are_content_addressed_and_append_only(seeded):
    a = intake.submit("EVIDENCE", b"LR No: 1\n", "lr.txt", "whs.ganesh")
    b = intake.submit("EVIDENCE", b"LR No: 1\n", "lr-copy.txt", "whs.ganesh")
    with db.borrow("pbc", "audit_staff") as p:
        rows = p.execute("SELECT sha256, stored_path FROM submissions WHERE id IN (?, ?)", (a, b)).fetchall()
    assert rows[0][0] == rows[1][0] and rows[0][1] == rows[1][1]          # same bytes, one stored file, two rows
    with pytest.raises(ValueError, match="empty"):
        intake.submit("EVIDENCE", b"", "x.txt", "whs.ganesh")
