"""AA-3 / AA-4 / DA-4: findings with balanced AJEs (06 §6), SA 450, the 4-sheet workpaper (Sheet 1 cell-for-cell,
07 §8.4), the memo, and the partner's decide / sign-off (06 §9)."""
import io
import json
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from app import analytics, config, db, erp, evidence, findings, intake, memo, sampling, seed, workpaper

ACC, WHS, SAL = ("acc.neha", "accounts"), ("whs.ganesh", "warehouse"), ("sales.priya", "sales")
PARTNER = {"principal": "ptn.mehta", "role": "engagement_partner", "name": "CA Rakesh Mehta", "membership_no": "118204"}
L = 1_00_000_00


@pytest.fixture()
def engaged(world):
    plan = seed.build_plan(120, 7, network=False)
    seed.load(plan, reset=True)
    P = {}

    def sale(qty=10, d="2026-03-30"):
        config.set_clock(f"{d}T10:00:00+05:30")
        so = erp.create_sales_order(*SAL, party_id="P-0003", confirmed_on="2026-03-01",
                                    lines=[{"item_id": "IT-001", "qty": qty}])["result"]["id"]
        return erp.issue_invoice(*ACC, so_id=so, invoice_date=d)["result"]
    P["premature"] = sale()
    config.set_clock("2026-04-03T10:00:00+05:30")
    erp.record_dispatch(*WHS, invoice_nos=[P["premature"]["id"]], moved_on="2026-04-03", lr_no="LR-P1")
    P["ghost"] = sale(qty=150)
    config.set_clock("2026-04-05T10:00:00+05:30")
    P["jv"] = erp.post_journal_entry(*ACC, voucher_date="2026-03-31", narration="year end adjustment for March",
                                     lines=[{"account_code": "SUSPENSE", "dr_paise": 12 * L},
                                            {"account_code": "SALES", "cr_paise": 12 * L}])["result"]
    config.set_clock("2026-10-05T10:00:00+05:30")
    intake.on_client_upload(intake.submit("RP_LIST", seed.declared_rp_csv(plan).encode(), "rp.csv", "adm.kiran"))
    # the client's register as an xlsx with a few awkward cells, exactly as a judge might send it
    exp = erp.export_sales_register(*ACC)["result"]
    with db.borrow("pbc", "audit_staff") as p:
        path = p.execute("SELECT stored_path FROM exports WHERE id = ?", (exp["export_id"],)).fetchone()[0]
    wb = load_workbook(path)
    ws = wb.active
    c = ws.cell(row=2, column=4)
    c.value = '=HYPERLINK("http://x","ok")'
    c.data_type = "s"
    ws.cell(row=3, column=21).value = "#N/A"
    ws.cell(row=3, column=21).data_type = "s"
    ws.cell(row=4, column=8).number_format = "0.000"
    buf = io.BytesIO()
    wb.save(buf)
    P["register_bytes"] = buf.getvalue()
    intake.on_client_upload(intake.submit("SALES_REGISTER", P["register_bytes"], "Sales Register FY26.xlsx", "acc.neha"))
    analytics.run_all(config.ENGAGEMENT_ID)
    sampling.select(config.ENGAGEMENT_ID, "sales_tod", actor="aud.judge")
    sampling.select(config.ENGAGEMENT_ID, "cutoff", actor="aud.judge")
    return P


def test_findings_have_impacts_standards_and_balanced_ajes(engaged):
    P = engaged
    new = findings.from_hits(config.ENGAGEMENT_ID)
    assert new and findings.from_hits(config.ENGAGEMENT_ID) == []                 # idempotent
    with db.borrow("audit", "audit_staff") as a:
        fs = {(r["detector"], json.loads(r["doc_refs_json"])[0]): dict(r) for r in a.execute("SELECT * FROM findings")}
    pre, ghost, jv = fs[("D3", P["premature"]["id"])], fs[("D4", P["ghost"]["id"])], fs[("D7", str(P["jv"]["journal_no"]))]
    for f in fs.values():
        assert f["standards"] and f["explanation"] and f["status"] == "PROPOSED" and f["confidence"] in ("HIGH", "MEDIUM", "LOW")
        if f["proposed_aje_json"]:
            lines = json.loads(f["proposed_aje_json"])
            assert sum(x["dr_paise"] for x in lines) == sum(x["cr_paise"] for x in lines) > 0
    a1 = {x["account_code"]: x for x in json.loads(pre["proposed_aje_json"])}
    assert a1["SALES"]["dr_paise"] == a1["CONTRACT_LIABILITY"]["cr_paise"] == P["premature"]["taxable_paise"]
    assert "INVENTORY" in a1 and "OUTPUT_CGST" not in a1                         # GST is not reversed (06 §6)
    a2 = {x["account_code"]: x for x in json.loads(ghost["proposed_aje_json"])}
    assert a2["DEBTORS"]["cr_paise"] == P["ghost"]["taxable_paise"] + P["ghost"]["tax_paise"]
    a3 = {x["account_code"]: x for x in json.loads(jv["proposed_aje_json"])}
    assert a3["SALES"]["dr_paise"] == a3["SUSPENSE"]["cr_paise"] == 12 * L
    assert all(f["proposed_aje_json"] is None for f in fs.values() if f["detector"] in ("D0", "D1", "D2"))
    s = findings.sa450_summary(config.ENGAGEMENT_ID)
    assert s["factual_paise"] >= P["premature"]["taxable_paise"] + 12 * L and s["materiality_paise"] > 0


def test_sheet1_is_the_uploaded_file_cell_for_cell(engaged):
    findings.from_hits(config.ENGAGEMENT_ID)
    out = workpaper.build(config.ENGAGEMENT_ID)
    wb = load_workbook(out["path"])
    assert wb.sheetnames == ["Sales Reg.", "Sales TOD", "Sales Cut-off TOD", "Credit Note TOD"]
    src = load_workbook(io.BytesIO(engaged["register_bytes"])).active
    ws = wb["Sales Reg."]
    for r in range(1, src.max_row + 1):
        for c in range(1, src.max_column + 1):
            a, b = src.cell(row=r, column=c), ws.cell(row=r + 5, column=c)
            assert a.value == b.value, (r, c, a.value, b.value)
            if a.value is not None:
                assert a.number_format == b.number_format, (r, c)
    assert ws.cell(row=7, column=4).data_type == "s" and ws.cell(row=7, column=4).value.startswith("=HYPERLINK")
    assert ws.cell(row=8, column=21).value == "#N/A" and ws.cell(row=8, column=21).data_type == "s"
    assert ws.cell(row=6, column=src.max_column + 1).value == "[Auditor] SoR_Match"
    assert ws["C1"].value.startswith("DRAFT") and "DRAFT" in ws.oddHeader.center.text
    tod = wb["Sales TOD"]
    assert tod["B4"].value == "Sales TOD" and tod.cell(row=18, column=1).value == "Sr No." and tod.cell(row=19, column=1).value == 1
    assert any("Sales Cut off TOD" == wb["Sales Cut-off TOD"]["B4"].value for _ in [0])


def test_memo_then_partner_sign_off(engaged):
    findings.from_hits(config.ENGAGEMENT_ID)
    m = memo.build(config.ENGAGEMENT_ID)
    md = open(m["path"], encoding="utf-8").read()
    assert "DRAFT – NOT REVIEWED" in md and all(f"## {i}." in md for i in range(1, 9)) and "SA 705" in md
    assert open(m["html_path"], encoding="utf-8").read().startswith("<!doctype html>")
    with pytest.raises(ValueError, match="still PROPOSED"):
        findings.sign_off(config.ENGAGEMENT_ID, PARTNER)
    with pytest.raises(PermissionError):
        findings.decide("F-001", "CONFIRMED", "ok", {**PARTNER, "role": "audit_staff"})
    with pytest.raises(ValueError, match="comment is mandatory"):
        findings.decide("F-001", "CONFIRMED", " ", PARTNER)
    with db.borrow("audit", "audit_staff") as a:
        ids = [r[0] for r in a.execute("SELECT id FROM findings ORDER BY id")]
    findings.decide(ids[0], "REJECTED", "explained by management: timing of POD scan", PARTNER)
    for fid in ids[1:]:
        findings.decide(fid, "CONFIRMED", "agreed with the evidence", PARTNER)
    assert evidence.verify_integrity(config.ENGAGEMENT_ID)["ok"]
    r = findings.sign_off(config.ENGAGEMENT_ID, PARTNER)
    assert len(r["signoff_hash"]) == 64 and r["anchor"]["label"] == "SIGNOFF"
    final_md = open(r["memo"]["path"], encoding="utf-8").read()
    assert "DRAFT" not in final_md and "Status FINAL" in final_md and "Rejected by the engagement partner" in final_md
    assert "ICAI M.No. 118204" in final_md
    wb = load_workbook(r["workpaper"]["path"])
    assert wb["Sales Reg."]["C1"].value is None
    with db.borrow("audit", "engagement_partner") as a:
        assert a.execute("SELECT status FROM engagements").fetchone()[0] == "SIGNED_OFF"
        with pytest.raises(Exception, match="signed off"):
            a.execute("UPDATE findings SET status = 'PROPOSED' WHERE id = ?", (ids[1],))
